#!/usr/bin/env python3
"""Renumber Ch3 figures 3.2..3.9 -> +1, then insert the data-flow diagram as the
new Figure 3.2 in section 3.1 (after the Path A/B paragraph) with an explanation."""
import re
from io import BytesIO
from docx import Document
from docx.shared import Cm
from PIL import Image

DOC = 'senior_project_report.docx'
IMG = 'figures/figure3.2_data_flow.png'

CAPTION = ("Figure 3.2  End-to-end data flow across the IT/OT stack, coloured by "
           "protocol: Modbus TCP (blue), WebSocket (green), and MQTT publish/subscribe "
           "(orange); the dashed path marks the monitoring flow into the IT stack.")

EXPLAIN = (
    "Figure 3.2 presents the same system from a protocol and data-flow perspective, "
    "complementing the physical layout of Figure 3.1. At the field edge, the ESP32-S3 "
    "(MCU) simultaneously exposes three interfaces: a Modbus TCP server, a WebSocket "
    "server for real-time streaming, and an MQTT publisher. These feed the Raspberry "
    "Pi 4 industrial gateway (SBC), which hosts MBPoll as a Modbus TCP client/tester, "
    "Node-RED for flow logic and orchestration, and the Mosquitto MQTT broker. Control "
    "traffic continues over Modbus TCP from both the MCU and the gateway to the Siemens "
    "LOGO! 8.4 PLC, which drives the physical I/O panel — relays and contactors, digital "
    "and analogue outputs, sensors, and motors and actuators. On the IT side (Data and "
    "Visualisation), Telegraf subscribes to the Mosquitto broker, persists the telemetry "
    "in the InfluxDB time-series database, and Grafana renders the dashboards. The colour "
    "coding separates the three transport protocols — Modbus TCP for industrial control, "
    "WebSocket for real-time communication, and MQTT for publish/subscribe messaging.")


def set_text_keep_style(p, text):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.add_run(text)


def main():
    doc = Document(DOC)

    # 1) renumber Figure 3.9..3.2 -> +1 (descending), skip front-matter list rows (tabs)
    for N in range(9, 1, -1):
        pat = re.compile(rf'Figure 3\.{N}\b')
        repl = f'Figure 3.{N + 1}'
        for p in doc.paragraphs:
            if '\t' in p.text:
                continue
            for r in p.runs:
                if r.text and f'3.{N}' in r.text:
                    r.text = pat.sub(repl, r.text)

    # 2) find anchor: the Path A/B paragraph in section 3.1
    anchor = None
    for p in doc.paragraphs:
        if p.text.strip().startswith('Path A carries'):
            anchor = p
            break
    if anchor is None:
        raise SystemExit('Path A/B anchor not found')
    anchor_el = anchor._p

    # 3) build image paragraph, caption, explanation; insert in order after anchor
    buf = BytesIO()
    Image.open(IMG).convert('RGB').save(buf, 'PNG')
    buf.seek(0)
    imgp = doc.add_paragraph()
    imgp.add_run().add_picture(buf, width=Cm(15.5))
    imgp.alignment = 1

    cap = doc.add_paragraph(CAPTION)
    expl = doc.add_paragraph(EXPLAIN)

    ref = anchor_el
    for para in (imgp, cap, expl):
        el = para._p
        el.getparent().remove(el)
        ref.addnext(el)
        ref = el

    doc.save(DOC)
    # report
    d2 = Document(DOC)
    caps = [p.text.strip()[:60] for p in d2.paragraphs
            if re.match(r'^Figure 3\.\d+\s\s', p.text.strip()) and '\t' not in p.text]
    print('Figure 3.x captions now:')
    for c in caps:
        print('  ', c)


if __name__ == '__main__':
    main()

#!/usr/bin/env python3
"""Insert Table 3.1 (Purdue level assignment) replacing the placeholder at para 271."""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm

DOC = 'senior_project_report.docx'

CAPTION = ('Table 3.1  Mapping of the testbed devices to the Purdue Enterprise '
           'Reference Architecture (ISA-95) levels.')

HEADERS = ['Purdue level', 'Device(s) in this testbed', 'Function', 'Protocol / interface']

ROWS = [
    ['L0 — Physical process (field)',
     'XY-MD02 temp/humidity sensor; SDM230 power meter; relay-driven pump/loads',
     'Sense physical quantities (temperature, humidity, voltage, current, power) and actuate field loads',
     'Modbus RTU over RS-485, 9,600 baud'],
    ['L1 — Basic control',
     'Siemens LOGO! 8.4 PLC; STM32F767ZI Nucleo (Study A test slave)',
     'Deterministic ladder control (coils Q1-Q4, V-memory HR0-HR9); STM32 provides a jitter-free Modbus test target',
     'Modbus TCP server, port 502'],
    ['L2 — Supervisory / edge',
     'Waveshare ESP32-S3-Relay-6CH gateway',
     'Protocol bridge OT<->IT: Modbus RTU master + Modbus TCP slave (HR0-HR20) + MQTT publisher',
     'RS-485 (south); Wi-Fi MQTT + Modbus TCP (north)'],
    ['L3 — Site operations',
     'Raspberry Pi 4 (Mosquitto, Telegraf, InfluxDB 2, Grafana) + 7" kiosk',
     'Broker MQTT traffic, ingest and store time-series, serve dashboards and visualisation',
     'MQTT (1883), HTTP/Flux API, Grafana UI'],
    ['L4-L5 — Enterprise / business',
     '-',
     'ERP, business logistics and cloud integration - out of scope for this project',
     '-'],
]

COL_W = [Cm(3.4), Cm(4.4), Cm(5.0), Cm(3.6)]


def set_cell(cell, text, bold=False):
    cell.text = ''
    r = cell.paragraphs[0].add_run(text)
    r.bold = bold


def main():
    doc = Document(DOC)

    anchor = None
    for p in doc.paragraphs:
        if p.text.strip().startswith('[Table 3.1'):
            anchor = p
            break
    if anchor is None:
        raise SystemExit('placeholder for Table 3.1 not found')
    anchor_el = anchor._p

    cap = doc.add_paragraph(CAPTION)
    cap_el = cap._p
    cap_el.getparent().remove(cap_el)
    anchor_el.addprevious(cap_el)

    table = doc.add_table(rows=1 + len(ROWS), cols=len(HEADERS))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for ci, h in enumerate(HEADERS):
        set_cell(table.rows[0].cells[ci], h, bold=True)
    for ri, row in enumerate(ROWS, start=1):
        for ci, val in enumerate(row):
            set_cell(table.rows[ri].cells[ci], val, bold=(ci == 0))

    for r in table.rows:
        for ci, c in enumerate(r.cells):
            c.width = COL_W[ci]

    tbl_el = table._tbl
    tbl_el.getparent().remove(tbl_el)
    anchor_el.addprevious(tbl_el)

    anchor_el.getparent().remove(anchor_el)

    doc.save(DOC)
    print('Inserted Table 3.1 with', len(ROWS), 'rows. Tables now:', len(doc.tables))


if __name__ == '__main__':
    main()

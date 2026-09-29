#!/usr/bin/env python3
"""Replace the Figure E.4 photo (self-assembled DIN-rail, plc.jpg) with the Siemens
manufacturer product image (SiemensLOGO.jpeg) and update the caption."""
from io import BytesIO
from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn
from PIL import Image

DOC = 'senior_project_report.docx'
IMG = 'figures/SiemensLOGO.jpeg'
CAPTION = ("Figure E.4  Siemens LOGO! 8.4 logic module (Siemens product image): eight "
           "24 V digital inputs (I1–I8, four usable as 0–10 V analogue), four outputs "
           "(Q1–Q4), integrated Ethernet (RJ-45), and a backlit display with six-button "
           "keypad.")


def istab(p):
    return p._p.find('.//' + qn('w:tab')) is not None


def main():
    doc = Document(DOC)
    cap = None
    for p in doc.paragraphs:
        if not istab(p) and p.text.strip().startswith('Figure E.4'):
            cap = p
            break
    if cap is None:
        raise SystemExit('Figure E.4 caption not found')

    old_img = cap._p.getprevious()
    if old_img is None or old_img.find('.//' + qn('w:drawing')) is None:
        raise SystemExit('preceding image paragraph not found')

    # build new centered image paragraph
    buf = BytesIO()
    Image.open(IMG).convert('RGB').save(buf, 'JPEG', quality=90)
    buf.seek(0)
    newp = doc.add_paragraph()
    newp.add_run().add_picture(buf, width=Cm(13.0))
    newp.alignment = 1
    # move new image before caption, then remove old image paragraph
    el = newp._p
    el.getparent().remove(el)
    cap._p.addprevious(el)
    old_img.getparent().remove(old_img)

    # update caption text (preserve style)
    for r in list(cap.runs):
        r._element.getparent().remove(r._element)
    cap.add_run(CAPTION)

    doc.save(DOC)
    print('Figure E.4 image swapped to manufacturer product photo; caption updated.')


if __name__ == '__main__':
    main()

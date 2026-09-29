#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Apply KMUTNB Thai typography (Angsana New) to thesis_th.docx, style-driven."""
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

import os
PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "thesis_th.docx")
FONT = "Angsana New"

d = docx.Document(PATH)

def is_code(r):
    rpr = r._element.find(qn('w:rPr'))
    if rpr is None:
        return False
    rf = rpr.find(qn('w:rFonts'))
    return rf is not None and (rf.get(qn('w:ascii')) or '').startswith(('Courier', 'Consolas'))

def set_run(r, size_pt, bold=None):
    if is_code(r):
        return
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), FONT)
    for tag in ('w:sz', 'w:szCs'):
        e = rpr.find(qn(tag))
        if e is None:
            e = OxmlElement(tag); rpr.append(e)
        e.set(qn('w:val'), str(int(size_pt * 2)))
    if bold is not None:
        for tag in ('w:b', 'w:bCs'):
            e = rpr.find(qn(tag))
            if e is None:
                e = OxmlElement(tag); rpr.append(e)
            e.set(qn('w:val'), '1' if bold else '0')

STYLE_RULES = {
    'Heading 1': (24, True, WD_ALIGN_PARAGRAPH.CENTER),
    'Heading 0': (18, True, WD_ALIGN_PARAGRAPH.CENTER),
    'Heading 2': (16, True, WD_ALIGN_PARAGRAPH.LEFT),
    'Heading 3': (16, True, WD_ALIGN_PARAGRAPH.LEFT),
    'Caption':   (16, None, WD_ALIGN_PARAGRAPH.LEFT),
    'Thesisname': (18, True, WD_ALIGN_PARAGRAPH.CENTER),
}
DEFAULT_RULE = (16, None, WD_ALIGN_PARAGRAPH.JUSTIFY)

count = 0
for p in d.paragraphs:
    size, bold, align = STYLE_RULES.get(p.style.name, DEFAULT_RULE)
    if align is not None:
        p.alignment = align
    for r in p.runs:
        if not r.text:
            continue
        # preserve explicit bold already set on the run (e.g. appendix A.1/A.2/A.3
        # sub-heading runs, table headers) unless the style forces a value
        run_bold = bold if bold is not None else r.bold
        set_run(r, size, bold=run_bold)
    count += 1

for t in d.tables:
    for ri, row in enumerate(t.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    if not r.text:
                        continue
                    run_bold = r.bold if r.bold is not None else (ri == 0)
                    set_run(r, 14, bold=run_bold)

# Normal style default font (affects any un-run paragraph marks / new content)
try:
    nm = d.styles['Normal']
    nm.font.name = FONT
    rpr = nm.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), FONT)
except Exception as e:
    print("Normal style skip:", e)

d.save(PATH)
print("Applied Thai typography to", count, "paragraphs and", len(d.tables), "tables.")

#!/usr/bin/env python3
"""Format the Thai chapter files using the ACTUAL KMUTNB template styles.
1) Inject the template's styles.xml (+ theme) into each docx so the predefined
   named styles exist (Normal, Heading0/1/2/3, caption, Abstract, ...).
2) Assign every paragraph to the correct named style and strip direct run/para
   formatting so the template styles fully govern (code runs are preserved)."""
import os, re, glob, zipfile, shutil
from docx import Document
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOTX_DIR = '/tmp/dotx'                      # unpacked template
SRC = '/tmp/thai_bak'                       # clean original Thai chapters
DST = 'chapter_thai'

CH  = re.compile(r'^บทที่\s*\d')
SEC = re.compile(r'^\d+\.\d+\s')
SUB = re.compile(r'^\d+\.\d+\.\d+')
CAP = re.compile(r'^(รูปที่|ภาพที่|ตารางที่)\s')
APX = re.compile(r'^(ภาคผนวก|APPENDIX)')
UNNUM = {'บทคัดย่อ', 'กิตติกรรมประกาศ', 'เอกสารอ้างอิง', 'ประวัติผู้แต่ง',
         'ประวัติผู้เขียน', 'สารบัญ', 'สารบัญภาพ', 'สารบัญตาราง', 'ศัพท์เฉพาะ'}
FIGCAP = ('รูปที่', 'ภาพที่')


def inject_styles(path):
    with open(os.path.join(DOTX_DIR, 'word/styles.xml'), 'rb') as f:
        styles = f.read()
    # strip style-level auto numbering/bullets (headings already carry manual
    # numbers like "บทที่ 3" / "3.1"); otherwise Word adds a "•" or "1." prefix.
    styles = re.sub(rb'<w:numPr>.*?</w:numPr>', b'', styles, flags=re.S)
    styles = re.sub(rb'<w:numPr\s*/>', b'', styles)
    theme = None
    tp = os.path.join(DOTX_DIR, 'word/theme/theme1.xml')
    if os.path.exists(tp):
        theme = open(tp, 'rb').read()
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        names = set(zin.namelist())
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == 'word/styles.xml':
                data = styles
            elif item.filename == 'word/theme/theme1.xml' and theme is not None:
                data = theme
            zout.writestr(item, data)
    os.replace(tmp, path)


def clear_run(r):
    rpr = r._element.find(qn('w:rPr'))
    if rpr is None:
        return
    for tag in ('w:rFonts', 'w:sz', 'w:szCs', 'w:b', 'w:bCs', 'w:i', 'w:iCs', 'w:color'):
        e = rpr.find(qn(tag))
        if e is not None:
            rpr.remove(e)


def is_code(r):
    rpr = r._element.find(qn('w:rPr'))
    if rpr is None:
        return False
    rf = rpr.find(qn('w:rFonts'))
    return rf is not None and (rf.get(qn('w:ascii')) or '').startswith('Courier')


def clear_para_direct(p):
    ppr = p._p.find(qn('w:pPr'))
    if ppr is None:
        return
    for tag in ('w:jc',):
        e = ppr.find(qn(tag))
        if e is not None:
            ppr.remove(e)


def prev_is_image(p):
    el = p._p.getprevious()
    while el is not None and el.tag == qn('w:p') and not (el.findtext('.//' + qn('w:t')) or '').strip() \
            and el.find('.//' + qn('w:drawing')) is None:
        el = el.getprevious()
    return el is not None and el.tag == qn('w:p') and el.find('.//' + qn('w:drawing')) is not None


def next_is_table(p):
    el = p._p.getnext()
    while el is not None and el.tag == qn('w:p') and not (el.findtext('.//' + qn('w:t')) or '').strip():
        el = el.getnext()
    return el is not None and el.tag == qn('w:tbl')


def has_image(p):
    return p._p.find('.//' + qn('w:drawing')) is not None


def main():
    os.makedirs(DST, exist_ok=True)
    for src in sorted(glob.glob(os.path.join(SRC, '*.docx'))):
        name = os.path.basename(src)
        path = os.path.join(DST, name)
        shutil.copy(src, path)
        inject_styles(path)

        d = Document(path)
        S = {s.name: s for s in d.styles}

        def style(*cands):
            for c in cands:
                if c in S:
                    return S[c]
            return None

        H0 = style('Heading 0', 'heading 0')
        H1 = style('heading 1', 'Heading 1')
        H2 = style('heading 2', 'Heading 2')
        H3 = style('heading 3', 'Heading 3')
        CAPS = style('caption', 'Caption')
        LP = style('List Paragraph', 'ListParagraph')
        NORM = style('Normal')
        ABS = style('Abstract')
        ACK = style('Acknowledgement')
        BIB = style('Bibliography')
        BIO = style('Biography')
        TOC1 = style('toc 1', 'TOC1')
        TOC2 = style('toc 2', 'TOC2')
        TOC3 = style('toc 3', 'TOC3')
        TOF = style('table of figures', 'TableofFigures')

        # fix mistranslation
        for p in d.paragraphs:
            for r in p.runs:
                if r.text and 'เชิงนามธรรม' in r.text:
                    r.text = r.text.replace('เชิงนามธรรม', 'บทคัดย่อ')

        mode = 'normal'                       # for Abstract/Ack/Bib/Bio body
        for p in d.paragraphs:
            t = p.text.strip()
            sty = None
            bullet = p.style.name.lower().startswith('list')

            # TOC / List-of-Figures / List-of-Tables entries carry a <w:tab/> +
            # page number -> assign TOC styles, NEVER Heading styles.
            has_tab = p._p.find('.//' + qn('w:tab')) is not None
            if has_tab and (CH.match(t) or SEC.match(t) or SUB.match(t)
                            or t.split('\t')[0].strip() in UNNUM or APX.match(t)
                            or CAP.match(t)):
                if t.startswith(FIGCAP) or t.startswith('ตารางที่'):
                    sty = TOF
                elif SUB.match(t):
                    sty = TOC3
                elif SEC.match(t):
                    sty = TOC2
                else:
                    sty = TOC1
                if sty is not None:
                    p.style = sty
                for r in p.runs:
                    if not is_code(r):
                        clear_run(r)
                continue

            if CH.match(t):
                sty, mode = H1, 'normal'
            elif APX.match(t) and len(t) < 80:
                sty, mode = H0, 'normal'
            elif t in UNNUM:
                sty = H0
                mode = {'บทคัดย่อ': 'abs', 'กิตติกรรมประกาศ': 'ack',
                        'เอกสารอ้างอิง': 'bib', 'ประวัติผู้แต่ง': 'bio',
                        'ประวัติผู้เขียน': 'bio'}.get(t, 'normal')
            elif SUB.match(t):
                sty, mode = H3, 'normal' if False else mode
                mode = 'normal'
            elif SEC.match(t):
                sty, mode = H2, 'normal'
            elif t and CAP.match(t) and t.startswith(FIGCAP) and prev_is_image(p):
                sty = CAPS
            elif t and CAP.match(t) and t.startswith('ตารางที่') and next_is_table(p):
                sty = CAPS
            elif bullet:
                sty = LP
            elif t.startswith('คำสำคัญ') or t.startswith('Keywords'):
                sty = ABS or NORM
            else:
                sty = {'abs': ABS, 'ack': ACK, 'bib': BIB, 'bio': BIO}.get(mode) or NORM

            if sty is not None:
                p.style = sty
            clear_para_direct(p)
            if has_image(p):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                if not is_code(r):
                    clear_run(r)

        # table cells -> Normal, strip direct formatting (keep code)
        for tb in d.tables:
            for row in tb.rows:
                for c in row.cells:
                    for p in c.paragraphs:
                        if NORM is not None and not p.style.name.lower().startswith('list'):
                            p.style = NORM
                        for r in p.runs:
                            if not is_code(r):
                                clear_run(r)

        d.save(path)
        print('styled', name)
    print('done')


if __name__ == '__main__':
    main()

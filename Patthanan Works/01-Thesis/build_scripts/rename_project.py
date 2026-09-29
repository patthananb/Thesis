#!/usr/bin/env python3
"""Apply the new project title and reframe Abstract / Chapter 1 / Chapter 5 around
'an Educational IIoT Platform Combining SBC, MCU and PLC'."""
from docx import Document
from docx.oxml.ns import qn

DOC = 'senior_project_report.docx'
NEW_TITLE = 'Design and Implementation of an Educational IIoT Platform Combining SBC, MCU and PLC'
COVER1 = 'DESIGN AND IMPLEMENTATION OF AN EDUCATIONAL IIoT PLATFORM'
COVER2 = 'COMBINING SBC, MCU AND PLC'

ABS_OLD1 = ('This project presents the design and implementation of IT/OT integration in IIoT '
            'applications, realised as a self-hosted testbed integrating three industrial '
            'edge-device classes —')
ABS_NEW1 = ('This project presents the design and implementation of an educational Industrial '
            'IoT (IIoT) platform that combines the three principal classes of industrial '
            'edge device —')
ABS_OLD2 = 'a Siemens LOGO! 8.4 PLC — into an open-source industrial IoT monitoring stack.'
ABS_NEW2 = ('a Siemens LOGO! 24CE PLC — into a single, low-cost, open-source teaching testbed '
            'that demonstrates end-to-end IT/OT integration.')
KEYWORDS = ('Keywords: educational IIoT platform, IT/OT convergence, SBC/MCU/PLC, Modbus TCP, '
            'MQTT, Raspberry Pi, ESP32, Grafana')

OBJ_OLD = 'Build a reproducible IT/OT convergence testbed integrating an ESP32-S3 MCU, a Raspberry Pi 4 SBC, and a Siemens LOGO! 8.4 PLC over Modbus TCP and MQTT.'
OBJ_NEW = ('Design and build an affordable, reproducible educational IIoT platform that combines '
           'an ESP32-S3 MCU, a Raspberry Pi 4 SBC, and a Siemens LOGO! 24CE PLC over Modbus TCP '
           'and MQTT as a hands-on teaching testbed.')

CONC_OLD = 'This project successfully designs and implements IT/OT integration in IIoT applications using an ESP32-S3 MCU, Raspberry Pi 4 SBC, and Siemens LOGO! 8.4 PLC.'
CONC_NEW = ('This project successfully designs and implements an educational IIoT platform that '
            'combines an ESP32-S3 MCU, a Raspberry Pi 4 SBC, and a Siemens LOGO! 24CE PLC into a '
            'single low-cost testbed for teaching end-to-end IT/OT integration.')

EDU_PARA = ('This project addresses that gap from an educational standpoint. Although students '
            'commonly study microcontrollers, single-board computers, and programmable logic '
            'controllers in isolation, they rarely encounter all three working together across '
            'the IT/OT boundary. The platform developed here combines an MCU, an SBC, and a PLC '
            'into one affordable, reproducible testbed so that the concepts of industrial IoT '
            'integration can be taught and demonstrated end to end.')
SCOPE_ADD = (' The platform is intended for laboratory and teaching use rather than certified '
             'industrial deployment.')


def first_fmt(p):
    if p.runs:
        r = p.runs[0]
        return (r.bold, r.italic, r.font.size, r.font.name)
    return (None, None, None, None)


def rebuild(p, text):
    bold, italic, size, name = first_fmt(p)
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = size
    if name:
        r.font.name = name


def repl_full(p, old, new):
    rebuild(p, p.text.replace(old, new))


def main():
    doc = Document(DOC)
    ps = doc.paragraphs

    # cover title (para with the all-caps IT/OT line)
    for p in ps:
        if 'DESIGN AND IMPLEMENTATION OF IT/OT INTEGRATION' in p.text:
            bold, italic, size, name = first_fmt(p)
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            r1 = p.add_run(COVER1); r1.bold = bold
            if size: r1.font.size = size
            if name: r1.font.name = name
            br = p.add_run(); br.add_break()
            r2 = p.add_run(COVER2); r2.bold = bold
            if size: r2.font.size = size
            if name: r2.font.name = name
            break

    for p in ps:
        t = p.text.strip()
        if t.startswith('Project Title:') and 'IT/OT Integration in IIoT' in t:
            repl_full(p, 'Design and Implementation of IT/OT Integration in IIoT Applications', NEW_TITLE)
        elif t.startswith('This project presents the design and implementation of IT/OT'):
            rebuild(p, p.text.replace(ABS_OLD1, ABS_NEW1).replace(ABS_OLD2, ABS_NEW2))
        elif t.startswith('Keywords:'):
            rebuild(p, KEYWORDS)
        elif t.startswith('Build a reproducible IT/OT convergence testbed'):
            rebuild(p, OBJ_NEW)
        elif t.startswith('This project successfully designs and implements IT/OT integration'):
            repl_full(p, CONC_OLD, CONC_NEW)
        elif t.startswith('In scope:') and 'cloud deployment' in t:
            rebuild(p, p.text.rstrip() + SCOPE_ADD)

    # insert educational motivation paragraph after the first 1.1 body paragraph
    for i, p in enumerate(ps):
        if p.text.strip().startswith('Industry 4.0 requires OT field data'):
            newp = doc.add_paragraph(EDU_PARA)
            newp.style = p.style
            el = newp._p
            el.getparent().remove(el)
            p._p.addnext(el)
            break

    doc.save(DOC)
    print('Title + framing updated.')


if __name__ == '__main__':
    main()

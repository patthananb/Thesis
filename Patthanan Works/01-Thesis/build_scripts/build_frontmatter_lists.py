#!/usr/bin/env python3
"""Populate TABLE OF CONTENTS, LIST OF FIGURES, LIST OF TABLES with real entries.
Pass 1: insert entries with dummy page numbers and write a sidecar JSON of the
search keys (in order) so a second pass can fill in calibrated page numbers."""
import json, re
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER

DOC = 'senior_project_report.docx'
SIDE = 'frontmatter_keys.json'
TABW = Cm(15.5)


def norm(s):
    return re.sub(r'\s+', ' ', s).strip()


def collect(doc):
    fm = {'ABSTRACT', 'ACKNOWLEDGEMENTS', 'NOMENCLATURE', 'LIST OF FIGURES',
          'LIST OF TABLES', 'REFERENCES', 'BIOGRAPHY'}
    toc, lof, lot = [], [], []
    capnum = re.compile(r'^(Table|Figure)\s+([0-9]+|[A-E])\.[0-9]+\s\s+(\S.*)$')
    secnum = re.compile(r'^[1-5]\.\d+(\.\d+)?\s+\S')
    for p in doc.paragraphs:
        raw = p.text.strip()
        t = norm(raw)
        # TOC entries
        if t in fm:
            toc.append({'disp': t, 'key': t, 'lvl': 0})
        elif t.startswith('CHAPTER '):
            m = re.match(r'(CHAPTER \d+)\s+(.*)', t)
            toc.append({'disp': f'{m.group(1)}  {m.group(2)}', 'key': m.group(1), 'lvl': 0})
        elif t.startswith('APPENDIX '):
            m = re.match(r'(APPENDIX [A-E])', t)
            toc.append({'disp': t, 'key': m.group(1), 'lvl': 0})
        elif secnum.match(t):
            lvl = 2 if re.match(r'^[1-5]\.\d+\.\d+', t) else 1
            toc.append({'disp': t, 'key': t, 'lvl': lvl})
        # caption entries
        m = capnum.match(raw)
        if m:
            kind, _, rest = m.group(1), m.group(2), m.group(3)
            num = raw.split()[1]
            short = re.split(r'[:.]', norm(rest))[0]
            short = ' '.join(short.split()[:9])
            disp = f'{kind} {num}  {short}'
            key = norm(rest)[:40]
            (lof if kind == 'Figure' else lot).append({'disp': disp, 'key': key})
    return toc, lof, lot


def make_entry(doc, disp, lvl=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if lvl:
        pf.left_indent = Cm(0.7 * lvl)
    pf.tab_stops.add_tab_stop(TABW, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    p.add_run(disp)
    p.add_run('\t')
    p.add_run('0')  # page placeholder, fixed in pass 2
    return p


def insert_after_header(doc, header_text, entries):
    # find header paragraph and the following placeholder paragraph
    paras = doc.paragraphs
    hidx = None
    for i, p in enumerate(paras):
        if p.text.strip() == header_text:
            hidx = i
            break
    if hidx is None:
        raise SystemExit(f'header not found: {header_text}')
    anchor = paras[hidx]._p
    # remove the '[Update in Word...]' placeholder right after, if present
    nxt = paras[hidx + 1]
    if nxt.text.strip().startswith('[Update in Word'):
        nxt._p.getparent().remove(nxt._p)
    # build entries, then move them in order after the header
    built = []
    for e in entries:
        built.append(make_entry(doc, e['disp'], e.get('lvl', 0)))
    ref = anchor
    for b in built:
        el = b._p
        el.getparent().remove(el)
        ref.addnext(el)
        ref = el


def main():
    doc = Document(DOC)
    toc, lof, lot = collect(doc)
    insert_after_header(doc, 'TABLE OF CONTENTS', toc)
    insert_after_header(doc, 'LIST OF FIGURES', lof)
    insert_after_header(doc, 'LIST OF TABLES', lot)
    doc.save(DOC)
    json.dump({'toc': [e['key'] for e in toc],
               'lof': [e['key'] for e in lof],
               'lot': [e['key'] for e in lot]},
              open(SIDE, 'w'), ensure_ascii=False, indent=1)
    print(f'Inserted TOC={len(toc)} LoF={len(lof)} LoT={len(lot)} entries.')


if __name__ == '__main__':
    main()

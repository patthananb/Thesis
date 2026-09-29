#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Split the combined Section 2.6 (TIG + MQTT) into two sections.

Before:  2.6 Time-Series Monitoring Stack (TIG) and MQTT
         2.7 Related Work and Research Gap
After:   2.6 Time-Series Monitoring Stack (TIG)
         2.7 MQTT Broker (Mosquitto)         <-- new
         2.8 Related Work and Research Gap    <-- renumbered

Applied to both editions. New paragraphs copy the run formatting of the
existing 2.6 heading/body so fonts and sizing match. Idempotent: if 2.6 no
longer says "... and MQTT", the script does nothing.

Run as a late pipeline step (after drop_purdue.py):
    python3 build_scripts/split_tig_mqtt.py
"""
import copy
import os
import docx
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

EDITIONS = {
    'thesis_en.docx': dict(
        head26_old="2.6  Time-Series Monitoring Stack (TIG) and MQTT",
        head26_new="2.6  Time-Series Monitoring Stack (TIG)",
        body_find=" The Mosquitto broker connects the edge to this stack over "
                  "MQTT. All four components",
        body_repl=" The three components",
        head27_new="2.7  MQTT Broker (Mosquitto)",
        body27="The Mosquitto broker connects the edge devices to the "
               "monitoring stack over MQTT, the publish-subscribe protocol "
               "described in Section 2.3.3. Edge devices publish telemetry to "
               "topics on the broker, and Telegraf's mqtt_consumer input "
               "subscribes to those topics to ingest the data. Like the TIG "
               "components, Mosquitto is lightweight, widely deployed, and "
               "distributed as a Docker image orchestrated with Docker Compose.",
        related_old="2.7  Related Work and Research Gap",
        related_new="2.8  Related Work and Research Gap",
    ),
    'thesis_th.docx': dict(
        head26_old="2.6  ชุดเฝ้าระวังข้อมูลอนุกรมเวลา (TIG) และ MQTT",
        head26_new="2.6  ชุดแดชบอร์ดข้อมูลอนุกรมเวลา (TIG)",
        body_find=" โบรกเกอร์ Mosquitto เชื่อมต่อชั้นขอบเข้ากับชุดนี้ผ่าน MQTT ทั้งสี่องค์ประกอบ",
        body_repl=" ทั้งสามองค์ประกอบ",
        head27_new="2.7  โบรกเกอร์ MQTT (Mosquitto)",
        body27="โบรกเกอร์ Mosquitto เชื่อมต่ออุปกรณ์ชั้นขอบเข้ากับชุดแดชบอร์ดผ่าน MQTT "
               "ซึ่งเป็นโพรโทคอลแบบเผยแพร่-สมัครสมาชิกที่อธิบายไว้ในหัวข้อ 2.3.3 "
               "อุปกรณ์ชั้นขอบเผยแพร่ข้อมูลโทรมาตรไปยังหัวข้อ (topic) บนโบรกเกอร์ "
               "และอินพุต mqtt_consumer ของ Telegraf สมัครสมาชิกหัวข้อเหล่านั้นเพื่อนำข้อมูลเข้า "
               "เช่นเดียวกับองค์ประกอบของชุด TIG นั้น Mosquitto มีน้ำหนักเบา ใช้งานแพร่หลาย "
               "และถูกเผยแพร่เป็นอิมเมจ Docker ที่จัดการด้วย Docker Compose",
        related_old="2.7  งานวิจัยที่เกี่ยวข้องและช่องว่างงานวิจัย",
        related_new="2.8  งานวิจัยที่เกี่ยวข้องและช่องว่างงานวิจัย",
    ),
}


def set_text_keep_fmt(p, text):
    """Set paragraph text, keeping the first run's formatting."""
    runs = p.runs
    if not runs:
        p.add_run(text); return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def new_para_after(anchor_el, doc, style_para, text):
    """Create a paragraph after anchor_el, cloning style + first-run rPr of
    style_para, with the given text."""
    p_el = copy.deepcopy(style_para._p)
    # strip existing runs / content from the clone, keep pPr
    for child in list(p_el):
        if child.tag == qn('w:r'):
            p_el.remove(child)
    anchor_el.addnext(p_el)
    para = docx.text.paragraph.Paragraph(p_el, doc)
    run = para.add_run(text)
    # copy run formatting from the source paragraph's first run
    src_runs = style_para.runs
    if src_runs is not None and len(src_runs):
        src_rpr = src_runs[0]._r.find(qn('w:rPr'))
        if src_rpr is not None:
            run._r.insert(0, copy.deepcopy(src_rpr))
    return para


def process(fname, cfg):
    path = os.path.join(ROOT, fname)
    d = docx.Document(path)
    paras = d.paragraphs

    # locate 2.6 heading
    h26 = next((p for p in paras if p.text.strip() == cfg['head26_old']), None)
    if h26 is None:
        print(f"{fname}: already split (no combined 2.6 heading); skipping.")
        return

    # rename 2.6 heading
    set_text_keep_fmt(h26, cfg['head26_new'])

    # the 2.6 body is the next paragraph
    body = None
    el = h26._p.getnext()
    while el is not None:
        if el.tag == qn('w:p'):
            body = docx.text.paragraph.Paragraph(el, d)
            break
        el = el.getnext()
    # trim the MQTT sentence out of the TIG body
    if cfg['body_find'] in body.text:
        set_text_keep_fmt(body, body.text.replace(cfg['body_find'], cfg['body_repl']))

    # insert new 2.7 body then heading (insert body first, then heading before it)
    new_body = new_para_after(body._p, d, body, cfg['body27'])
    new_head = new_para_after(body._p, d, h26, cfg['head27_new'])
    # order is now: body, new_head, new_body  (addnext puts each right after body,
    # so the later insert ends up first) -> verify/fix ordering
    # body._p.next should be new_head, then new_body
    if new_head._p.getnext() is not new_body._p:
        # move new_body to directly after new_head
        new_body._p.getparent().remove(new_body._p)
        new_head._p.addnext(new_body._p)

    # renumber the old 2.7 (Related Work) -> 2.8
    rel = next((p for p in d.paragraphs if p.text.strip() == cfg['related_old']), None)
    if rel is not None:
        set_text_keep_fmt(rel, cfg['related_new'])

    d.save(path)
    print(f"{fname}: split 2.6 -> 2.6 (TIG) + 2.7 (MQTT); Related Work -> 2.8.")


if __name__ == '__main__':
    for fname, cfg in EDITIONS.items():
        if os.path.exists(os.path.join(ROOT, fname)):
            process(fname, cfg)

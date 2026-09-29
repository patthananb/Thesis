#!/usr/bin/env python3
"""Append APPENDIX F — IIoT Platform Comparison Matrix to the report, built from
iiot_comparison.xlsx, with category bands and cells colour-coded per the legend
(star=green, half=amber, cross=red)."""
import openpyxl
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC = 'senior_project_report.docx'
XLSX = 'iiot_comparison.xlsx'

NAVY = '0F2A43'
GREEN, AMBER, REDC = 'E3F2E7', 'FBEFD6', 'F7E1E1'
HEADERS = ['Dimension', 'ESP32-WROOM\n(MCU)', 'STM32 F7xx\n(MCU)', 'Arduino Mega\n(MCU)',
           'Raspberry Pi 4\n(SBC)', 'LOGO! 24CE\n(PLC)']
COLW = [Cm(3.0), Cm(2.5), Cm(2.5), Cm(2.5), Cm(2.5), Cm(2.5)]


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def set_cell(cell, text, bold=False, white=False, size=8, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if white:
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def color_for(text):
    t = text.strip()
    if t.startswith('★'):   # star
        return GREEN
    if t.startswith('◑'):   # half circle
        return AMBER
    if t.startswith('✗'):   # cross
        return REDC
    return None


def cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn('w:cantSplit')) is None:
        trPr.append(OxmlElement('w:cantSplit'))


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:tblHeader'); e.set(qn('w:val'), 'true'); trPr.append(e)


def load_rows():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb['IIoT Comparison']
    rows = list(ws.iter_rows(values_only=True))
    out = []
    for r in rows[2:]:                       # skip the 2 header rows
        cells = ['' if c is None else ' '.join(str(c).split()) for c in r[:6]]
        if not any(cells):
            continue
        is_cat = cells[0] and not any(cells[1:])
        out.append((is_cat, cells))
    return out


def main():
    doc = Document(DOC)

    # --- heading + intro (appended at end => after Appendix E) ---
    h = doc.add_paragraph()
    hr = h.add_run('APPENDIX F — IIoT PLATFORM COMPARISON MATRIX')
    hr.bold = True

    intro = doc.add_paragraph(
        'Table F.1 compares the candidate embedded platforms across six dimensions — '
        'industrial connectivity, wireless and networking, real-time and reliability, '
        'sensor and field I/O, data and edge compute, and deployment and maintenance. '
        'Three microcontroller-class devices (ESP32-WROOM, STM32 F7xx, Arduino Mega) are '
        'placed alongside the single-board computer (Raspberry Pi 4) and the PLC (Siemens '
        'LOGO! 24CE) used in this project. Cells are rated strong (★), partial (◑), '
        'or unsupported (✗), and shaded accordingly.')

    cap = doc.add_paragraph('Table F.1  IIoT platform comparison matrix (MCU vs SBC vs PLC).')

    data = load_rows()
    table = doc.add_table(rows=1 + len(data), cols=6)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # header
    for ci, htext in enumerate(HEADERS):
        c = table.rows[0].cells[ci]
        shade(c, NAVY)
        set_cell(c, htext, bold=True, white=True, size=8.5,
                 align=WD_ALIGN_PARAGRAPH.CENTER if ci else None)
    repeat_header(table.rows[0])

    # body
    for ri, (is_cat, cells) in enumerate(data, start=1):
        row = table.rows[ri]
        if is_cat:
            merged = row.cells[0]
            for ci in range(1, 6):
                merged = merged.merge(row.cells[ci])
            shade(merged, '33546E')
            set_cell(merged, cells[0], bold=True, white=True, size=8.5)
        else:
            for ci in range(6):
                c = row.cells[ci]
                set_cell(c, cells[ci], bold=(ci == 0), size=8)
                col = color_for(cells[ci]) if ci else None
                if col:
                    shade(c, col)
        cant_split(row)

    for row in table.rows:
        for ci, c in enumerate(row.cells):
            c.width = COLW[ci]

    # legend
    leg = doc.add_paragraph()
    for sym, txt, col in [('★', ' strong support / clear advantage   ', '2E7D32'),
                          ('◑', ' partial / conditional support   ', 'B7791F'),
                          ('✗', ' not supported / significant weakness', 'B23A48')]:
        r1 = leg.add_run(sym); r1.bold = True; r1.font.size = Pt(9)
        r1.font.color.rgb = RGBColor.from_string(col)
        r2 = leg.add_run(txt); r2.font.size = Pt(9)

    doc.save(DOC)
    print('Appendix F added with', len(data), 'rows (incl. category bands).')


if __name__ == '__main__':
    main()

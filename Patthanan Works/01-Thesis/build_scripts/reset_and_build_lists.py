#!/usr/bin/env python3
"""Clear the existing TOC/LoF/LoT entries and rebuild them from the current document
(picks up the new Figure 3.2 and renumbered figures). Writes sidecar keys for the
page-number calibration pass. Page numbers are inserted as '0' placeholders."""
import json, re
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER

DOC = 'senior_project_report.docx'
SIDE = 'frontmatter_keys.json'
TABW = Cm(15.5)
PLACE = '[Update in Word: References -> Update Table]'
ORDER = ['TABLE OF CONTENTS', 'LIST OF FIGURES', 'LIST OF TABLES', 'NOMENCLATURE']


def norm(s):
    return re.sub(r'\s+', ' ', s).strip()


def reset(doc):
    paras = doc.paragraphs
    idx = {h: next(i for i, p in enumerate(paras) if p.text.strip() == h) for h in ORDER}
    # remove paragraphs strictly between consecutive markers, then add one placeholder
    for a, b in [('TABLE OF CONTENTS', 'LIST OF FIGURES'),
                 ('LIST OF FIGURES', 'LIST OF TABLES'),
                 ('LIST OF TABLES', 'NOMENCLATURE')]:
        ha, hb = paras[idx[a]], paras[idx[b]]
        # collect elements strictly between
        to_del = []
        el = ha._p.getnext()
        while el is not None and el is not hb._p:
            to_del.append(el)
            el = el.getnext()
        for e in to_del:
            e.getparent().remove(e)
        # insert single placeholder after header a
        ph = doc.add_paragraph(PLACE)
        ph_el = ph._p
        ph_el.getparent().remove(ph_el)
        ha._p.addnext(ph_el)


def collect(doc):
    fm = {'ABSTRACT', 'ACKNOWLEDGEMENTS', 'NOMENCLATURE', 'LIST OF FIGURES',
          'LIST OF TABLES', 'REFERENCES', 'BIOGRAPHY'}
    toc, lof, lot = [], [], []
    capnum = re.compile(r'^(Table|Figure)\s+([0-9]+|[A-F])\.[0-9]+\s\s+(\S.*)$')
    secnum = re.compile(r'^[1-5]\.\d+(\.\d+)?\s+\S')

    def short(rest):
        s = norm(rest); w = s.split(); out = ' '.join(w[:10])
        if len(out) > 72:
            out = out[:72].rsplit(' ', 1)[0]
        return out.rstrip(' :.,;-')

    for p in doc.paragraphs:
        if '\t' in p.text:
            continue
        raw = p.text.strip(); t = norm(raw)
        if t in fm:
            toc.append({'disp': t, 'key': t, 'lvl': 0})
        elif t.startswith('CHAPTER '):
            m = re.match(r'(CHAPTER \d+)\s+(.*)', t)
            toc.append({'disp': f'{m.group(1)}  {m.group(2)}', 'key': m.group(1), 'lvl': 0})
        elif t.startswith('APPENDIX '):
            m = re.match(r'(APPENDIX [A-F])', t)
            toc.append({'disp': t, 'key': m.group(1), 'lvl': 0})
        elif secnum.match(t):
            lvl = 2 if re.match(r'^[1-5]\.\d+\.\d+', t) else 1
            toc.append({'disp': t, 'key': t, 'lvl': lvl})
        m = capnum.match(raw)
        if m:
            kind = m.group(1); num = raw.split()[1]
            disp = f'{kind} {num}  {short(m.group(3))}'
            key = norm(m.group(3))[:40]
            (lof if kind == 'Figure' else lot).append({'disp': disp, 'key': key})
    return toc, lof, lot


def make_entry(doc, disp, lvl=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if lvl:
        pf.left_indent = Cm(0.7 * lvl)
    pf.tab_stops.add_tab_stop(TABW, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    p.add_run(disp); p.add_run('\t'); p.add_run('0')
    return p


def insert_after_header(doc, header_text, entries):
    paras = doc.paragraphs
    hidx = next(i for i, p in enumerate(paras) if p.text.strip() == header_text)
    anchor = paras[hidx]._p
    nxt = paras[hidx + 1]
    if nxt.text.strip().startswith('[Update in Word'):
        nxt._p.getparent().remove(nxt._p)
    ref = anchor
    for e in entries:
        b = make_entry(doc, e['disp'], e.get('lvl', 0))
        el = b._p; el.getparent().remove(el); ref.addnext(el); ref = el


def main():
    doc = Document(DOC)
    reset(doc)
    toc, lof, lot = collect(doc)
    insert_after_header(doc, 'TABLE OF CONTENTS', toc)
    insert_after_header(doc, 'LIST OF FIGURES', lof)
    insert_after_header(doc, 'LIST OF TABLES', lot)
    doc.save(DOC)
    json.dump({'toc': [e['key'] for e in toc], 'lof': [e['key'] for e in lof],
               'lot': [e['key'] for e in lot]}, open(SIDE, 'w'), ensure_ascii=False, indent=1)
    print(f'Rebuilt TOC={len(toc)} LoF={len(lof)} LoT={len(lot)}')


if __name__ == '__main__':
    main()

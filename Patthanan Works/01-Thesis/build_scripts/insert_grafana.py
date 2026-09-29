#!/usr/bin/env python3
"""Add three Grafana dashboard screenshots to section 3.5:
  - Figure 3.9  Temperature & Humidity (TempHum.png)   [NEW, after 3.8 kiosk para]
  - Figure 3.10 SDM230 Power Meter (Powermeter.png)    [fills old 3.9 placeholder]
  - Figure 3.11 Pi host telemetry (PiTelemetry.png)    [NEW, after 3.10]
The existing IoT-Sniffer figure 3.10 is renumbered to 3.12.
"""
import re
from io import BytesIO
from docx import Document
from docx.shared import Cm
from PIL import Image

DOC = 'senior_project_report.docx'
WIDTH = Cm(14.0)

CAP = {
    '3.9': ("Figure 3.9  Grafana Temperature & Humidity dashboard (XY-MD02): temperature "
            "and relative-humidity time series over 30 days (last 35.2 °C / 50.4 %RH) "
            "with a sensor-reading status panel."),
    '3.10': ("Figure 3.10  Grafana SDM230 Power Meter dashboard: live single-phase readings "
             "— 229.1 V, 249 mA, 26.2 W active power — with power-over-time, power "
             "factor, frequency, and meter-status panels."),
    '3.11': ("Figure 3.11  Grafana Raspberry Pi host-telemetry dashboard: CPU temperature "
             "(61.3 °C), CPU load (14.7 %), and memory usage (82.0 %) with time-series "
             "history."),
}
IMGFILE = {'3.9': 'figures/TempHum.png', '3.10': 'figures/Powermeter.png',
           '3.11': 'figures/PiTelemetry.png'}


def image_par(doc, path):
    buf = BytesIO()
    Image.open(path).convert('RGB').save(buf, 'PNG')
    buf.seek(0)
    p = doc.add_paragraph()
    p.add_run().add_picture(buf, width=WIDTH)
    p.alignment = 1
    return p


def move_after(ref_el, paras):
    ref = ref_el
    for para in paras:
        el = para._p
        el.getparent().remove(el)
        ref.addnext(el)
        ref = el
    return ref


def main():
    doc = Document(DOC)

    # 1) renumber sniffer 3.10 -> 3.12 (captions + in-text refs), skip list rows
    for p in doc.paragraphs:
        if '\t' in p.text:
            continue
        for r in p.runs:
            if r.text and '3.10' in r.text:
                r.text = re.sub(r'Figure 3\.10\b', 'Figure 3.12', r.text)

    # 2) reference the figures in the narrative paragraph
    for p in doc.paragraphs:
        if p.text.startswith('Grafana 13.0.1 serves four dashboards') and 'Figures 3.9' not in p.text:
            p.runs[-1].text = p.runs[-1].text.rstrip()
            p.add_run(" The first three dashboards are shown in Figures 3.9–3.11.")
            break

    # 3) TempHum (3.9): insert after the kiosk paragraph
    kiosk = next(p for p in doc.paragraphs if p.text.strip().startswith('Kiosk: Pi OS Lite'))
    move_after(kiosk._p, [image_par(doc, IMGFILE['3.9']), doc.add_paragraph(CAP['3.9'])])

    # 4) SDM230 (3.10): replace the old "[Figure 3.9 ... SDM230 ...]" placeholder
    ph = next(p for p in doc.paragraphs if p.text.strip().startswith('[Figure 3.9'))
    ph_el = ph._p
    last = move_after(ph_el, [image_par(doc, IMGFILE['3.10']), doc.add_paragraph(CAP['3.10'])])
    ph_el.getparent().remove(ph_el)

    # 5) PiTelemetry (3.11): insert right after the 3.10 caption
    move_after(last, [image_par(doc, IMGFILE['3.11']), doc.add_paragraph(CAP['3.11'])])

    doc.save(DOC)
    d2 = Document(DOC)
    caps = [p.text.strip()[:58] for p in d2.paragraphs
            if re.match(r'^Figure 3\.\d+\s\s', p.text.strip()) and '\t' not in p.text]
    print('Figure 3.x captions now:')
    for c in caps:
        print('  ', c)


if __name__ == '__main__':
    main()

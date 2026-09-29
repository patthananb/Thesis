#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Restructure chapters 3, 4, 5 of thesis_th.docx (Thai edition only).

Chapter 3: rename 3.1/3.2, reframe 3.3-3.6 as software sections.
Chapter 4: rename 4.1 (drop "การศึกษา A:"), delete Study B (4.3) and table 4.2,
           keep figures 4.4/4.5 by moving them into 4.2.
Chapter 5: rewrite into 5.1 ภาพรวมระบบ / 5.2 ปัญหา อุปสรรค และวิธีการแก้ไข /
           5.3 ข้อเสนอแนะ (old 5.2-5.4 dropped, 5.5 -> 5.3).
Global:    เกตเวย์ขอบ -> เอดจ์เกตเวย์.

Note: §5.2 body is an AUTHOR-REVIEW DRAFT (marked inline).

Run once on the current thesis_th.docx (Thai-only, late step):
    python3 build_scripts/restructure_ch345_th.py
"""
import copy
import os
import docx
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOC = os.path.join(ROOT, "thesis_th.docx")

RENAME = {
    "3.1  สถาปัตยกรรมโดยรวมของแพลตฟอร์ม": "3.1  สถาปัตยกรรมโดยรวม",
    "3.2  การเลือกฮาร์ดแวร์และเหตุผลประกอบ": "3.2  การเลือกฮาร์ดแวร์ MCU, SBC และ PLC",
    "3.3  ชั้นโอที: PLC Siemens LOGO! และ LOGO! Soft Comfort":
        "3.3  ซอฟต์แวร์ชั้นโอที: การเขียนโปรแกรม PLC ด้วย LOGO! Soft Comfort",
    "3.4  เกตเวย์ขอบ: เฟิร์มแวร์ Modbus บน ESP32-S3":
        "3.4  เฟิร์มแวร์เอดจ์เกตเวย์: Modbus บน ESP32-S3",
    "3.5  ชั้นไอที: MQTT และชุด TIG แบบคอนเทนเนอร์":
        "3.5  ซอฟต์แวร์ชั้นไอที: MQTT และชุด TIG แบบคอนเทนเนอร์",
    "3.6  เครื่องมือดักจับแพ็กเก็ตไอโอที (IoT-Sniffer)": "3.6  เครื่องมือ IoT-Sniffer",
    "4.1  การศึกษา A: การประเมินไลบรารีเกตเวย์ Modbus":
        "4.1  การประเมินไลบรารีเกตเวย์ Modbus",
}

CH3_INTRO = ("บทนี้นำเสนอการออกแบบและพัฒนาแพลตฟอร์ม โดยเริ่มจากสถาปัตยกรรมโดยรวมและการนิยาม"
             "เส้นทางข้อมูล จากนั้นอธิบายการเลือกฮาร์ดแวร์ทั้งสามประเภท และสุดท้ายลงรายละเอียด"
             "ซอฟต์แวร์ของแต่ละชั้น ได้แก่ การเขียนโปรแกรม PLC เฟิร์มแวร์เอดจ์เกตเวย์บน ESP32-S3 "
             "ชุดไอทีแบบคอนเทนเนอร์ และเครื่องมือ IoT-Sniffer")

CH4_INTRO = ("บทนี้รายงานการประเมินสองส่วน ได้แก่ การทดสอบเชิงประจักษ์ของไลบรารีเกตเวย์ Modbus "
             "โอเพนซอร์สสำหรับ ESP32 และการวัดสมรรถนะด้านทรัพยากรและเวลาของเกตเวย์เชิงผลิตที่ได้")

CH5_INTRO = ("บทนี้สรุปภาพรวมของระบบที่พัฒนาขึ้น อภิปรายปัญหาและอุปสรรคที่พบระหว่างการดำเนินงาน"
             "พร้อมวิธีการแก้ไข และเสนอข้อเสนอแนะสำหรับการพัฒนาต่อยอดในอนาคต")

CH5_1_HEAD = "5.1  ภาพรวมระบบ"
CH5_1_BODY = ("แพลตฟอร์มที่พัฒนาขึ้นบูรณาการตัวควบคุมสามประเภทเข้าด้วยกันบนเครือข่ายท้องถิ่นเดียว "
              "ไมโครคอนโทรลเลอร์ ESP32-S3 ทำหน้าที่เอดจ์เกตเวย์แปลงโพรโทคอลจาก Modbus RTU ไปสู่ "
              "Modbus TCP และ MQTT คอมพิวเตอร์บอร์ดเดี่ยว Raspberry Pi 4 ทำหน้าที่ชั้นไอทีที่รัน"
              "โบรกเกอร์ MQTT และชุด TIG ในคอนเทนเนอร์ และพีแอลซี Siemens LOGO! 24CE ทำหน้าที่"
              "ควบคุมที่มีความแน่นอนเชิงเวลา ข้อมูลจากเครื่องมือวัดหน้างานไหลผ่านเอดจ์เกตเวย์ขึ้นสู่"
              "แดชบอร์ด Grafana ได้แบบเรียลไทม์ ผลการทดลองยืนยันว่าไมโครคอนโทรลเลอร์ราคาต่ำเพียง"
              "ตัวเดียวสามารถเชื่อมระบบโอทีอนุกรมเดิมเข้ากับแดชบอร์ดไอทีสมัยใหม่ได้ภายในงบประมาณ"
              "เฟิร์มแวร์ ตารางที่ 5.1 สรุปคุณลักษณะของตัวควบคุมทั้งสามประเภทเทียบกัน แสดงให้เห็นว่า"
              "อุปกรณ์แต่ละประเภทเหมาะสมกับบทบาทที่ต่างกันและเสริมกันมากกว่าจะแข่งขันกัน")

CH5_2_HEAD = "5.2  ปัญหา อุปสรรค และวิธีการแก้ไข"
CH5_2_BODY = ("(ส่วนนี้เป็นร่างที่เรียบเรียงจากเนื้อหาโครงงาน โปรดตรวจทานและปรับให้ตรงกับประสบการณ์จริง) "
              "ระหว่างการออกแบบและพัฒนาพบปัญหาและอุปสรรคหลายประการ พร้อมแนวทางแก้ไขดังนี้ ประการแรก "
              "ไลบรารีเกตเวย์ Modbus โอเพนซอร์สสำหรับ ESP32 มีให้เลือกจำนวนมากและมีคุณภาพแตกต่างกัน"
              "ทั้งด้านการบำรุงรักษาและความครอบคลุมคุณสมบัติ จึงยากต่อการตัดสินใจเลือก แก้ไขโดยกำหนด"
              "เกณฑ์การประเมินที่ชัดเจนและทดสอบสมรรถนะบนฮาร์ดแวร์จริงเพื่อคัดเลือกอย่างเป็นระบบ "
              "ประการที่สอง การสื่อสารผ่านบัส RS-485 มีความไวต่อการเดินสาย การสลับขั้ว A/B และการต่อ"
              "ความต้านทานปลายสาย ทำให้เกิดข้อผิดพลาด CRC เป็นระยะ แก้ไขโดยตรวจสอบการเดินสาย การต่อ"
              "กราวด์ร่วม และการตั้งค่า baud rate ให้ตรงกันทุกอุปกรณ์ ประการที่สาม หน่วยความจำแฟลชและแรม"
              "ของไมโครคอนโทรลเลอร์มีจำกัด เมื่อรวมหลายบทบาทไว้ในเฟิร์มแวร์เดียวจึงต้องบริหารทรัพยากร"
              "อย่างระมัดระวัง แก้ไขโดยเลือกไลบรารีที่ใช้ทรัพยากรน้อยและตรวจวัดการใช้หน่วยความจำอย่าง"
              "ต่อเนื่อง ประการที่สี่ การตั้งค่าชั้นไอทีที่ประกอบด้วยหลายบริการ (Mosquitto, InfluxDB, "
              "Telegraf และ Grafana) มีความซับซ้อนและยากต่อการทำซ้ำ แก้ไขโดยใช้ Docker Compose "
              "ประกาศค่าทั้งหมดเป็นโค้ดเพื่อให้ติดตั้งซ้ำได้อย่างสม่ำเสมอ ประการที่ห้า ภาพถ่ายฮาร์ดแวร์"
              "จากกล้องโทรศัพท์มีการหมุนตามข้อมูล EXIF ทำให้แสดงผลผิดทิศทางเมื่อนำเข้าเอกสาร แก้ไข"
              "โดยปรับทิศทางภาพให้ถูกต้องก่อนแทรกลงในรายงาน")

CH5_3_HEAD = "5.3  ข้อเสนอแนะ"
CH5_3_BODY = ("งานนี้สามารถพัฒนาต่อยอดได้หลายแนวทาง ด้านเอดจ์เกตเวย์ ควรเพิ่มบัฟเฟอร์วงแหวนสำหรับ"
              "ข้อมูลโทรมาตรเพื่อรองรับช่วงที่ Wi-Fi ขาดหาย และเพิ่มกลไกอัปเดตเฟิร์มแวร์ผ่านอากาศ "
              "(OTA) เพื่อความสะดวกในการบำรุงรักษา ด้านชั้นไอที ควรเพิ่มการสำรองข้อมูล InfluxDB ตาม"
              "กำหนดเวลา การตรวจสุขภาพคอนเทนเนอร์ และการแจ้งเตือนเมื่อค่าถึงขีดจำกัดใน Grafana เพื่อ"
              "เสริมความน่าเชื่อถือของชั้นข้อมูล ด้านการสื่อสาร ควรศึกษาการเพิ่มโพรโทคอล OPC-UA เพื่อ"
              "ขยายความสามารถในการทำงานร่วมกับระบบอุตสาหกรรมอื่น และด้านการนำไปใช้ ควรนำแพลตฟอร์ม"
              "ไปทดลองใช้จริงในรายวิชาปฏิบัติการร่วมกับนักศึกษาเพื่อประเมินคุณค่าทางการศึกษา ตลอดจน"
              "ต่อยอดสู่การใช้งานจริงหรืองานวิจัยในระดับที่สูงขึ้นต่อไป")

OBJ2_DROP = " และเปรียบเทียบไปป์ไลน์รับข้อมูลโอเพนซอร์สสำหรับชั้นไอที"


def txt(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t'))).strip()


def set_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text); return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def find_para(doc, pred):
    # skip TOC / List-of-Figures field-cache paragraphs (styled toc*/TableofFigures)
    for p in doc.paragraphs:
        sn = p.style.name.lower()
        if sn.startswith('toc') or sn == 'tableoffigures':
            continue
        if pred(p):
            return p
    return None


def new_para_after(anchor_el, doc, style_para, text):
    p_el = copy.deepcopy(style_para._p)
    for child in list(p_el):
        if child.tag == qn('w:r'):
            p_el.remove(child)
    anchor_el.addnext(p_el)
    para = docx.text.paragraph.Paragraph(p_el, doc)
    run = para.add_run(text)
    src = style_para.runs
    if src:
        rpr = src[0]._r.find(qn('w:rPr'))
        if rpr is not None:
            run._r.insert(0, copy.deepcopy(rpr))
    return para


def delete_block(start_para, stop_startswith, doc):
    el = start_para._p
    while el is not None:
        nxt = el.getnext()
        if el.tag == qn('w:p') and any(txt(el).startswith(s) for s in stop_startswith):
            break
        el.getparent().remove(el)
        el = nxt


def main():
    d = docx.Document(DOC)

    # -- heading renames
    for p in d.paragraphs:
        t = p.text.strip()
        if t in RENAME:
            set_text(p, RENAME[t])

    # -- chapter 3 & 4 intros
    p = find_para(d, lambda p: p.text.strip().startswith("บทนี้นำเสนอแพลตฟอร์มตั้งแต่ระดับสถาปัตยกรรม"))
    if p: set_text(p, CH3_INTRO)
    p = find_para(d, lambda p: p.text.strip().startswith("บทนี้รายงานการประเมินสามส่วน"))
    if p: set_text(p, CH4_INTRO)

    # -- objective 2 trim
    p = find_para(d, lambda p: p.text.strip().startswith("2)  สาธิตการบูรณาการ"))
    if p and OBJ2_DROP in p.text:
        set_text(p, p.text.replace(OBJ2_DROP, ""))

    # -- chapter 4: move figures 4.4/4.5 into 4.2, then delete Study B block
    def cap(prefix):
        return find_para(d, lambda p: p.style.name == 'Caption' and p.text.strip().startswith(prefix))
    cap43 = cap("ภาพที่ 4.3")
    c44, c45 = cap("ภาพที่ 4.4"), cap("ภาพที่ 4.5")
    if cap43 and c44 and c45:
        i44, i45 = c44._p.getprevious(), c45._p.getprevious()
        anchor = cap43._p
        for img, capel in ((i44, c44._p), (i45, c45._p)):
            anchor.addnext(capel); anchor.addnext(img); anchor = capel
    head43 = find_para(d, lambda p: p.text.strip().startswith("4.3  การศึกษา B"))
    if head43:
        delete_block(head43, ["บทที่ 5"], d)
    # cite the relocated figures 4.4/4.5 in the 4.2 body
    b42 = find_para(d, lambda p: p.style.name == 'Normal'
                    and p.text.strip().endswith("ดังแสดงในภาพที่ 4.2 และภาพที่ 4.3"))
    if b42:
        set_text(b42, b42.text + " นอกจากนี้ ภาพที่ 4.4 และภาพที่ 4.5 แสดงระบบขณะ"
                 "ทำงานจริง ได้แก่ แดชบอร์ดเฝ้าระวังทรัพยากรของ Raspberry Pi และจอสัมผัสที่"
                 "แสดงผล Grafana ในโหมดคีออสก์")

    # -- chapter 5 rewrite
    p = find_para(d, lambda p: p.text.strip().startswith("บทนี้ตีความผลการทดลอง"))
    if p: set_text(p, CH5_INTRO)
    h51 = find_para(d, lambda p: p.text.strip().startswith("5.1  MCU เทียบกับ SBC"))
    if h51: set_text(h51, CH5_1_HEAD)
    b51 = find_para(d, lambda p: p.text.strip().startswith("ข้อเสนอหลักในการออกแบบคือ"))
    if b51: set_text(b51, CH5_1_BODY)
    ref_h2 = h51  # style reference for new headings
    roles = find_para(d, lambda p: p.text.strip().startswith("MCU ให้ต้นทุนต่อช่องสัญญาณ"))

    # delete old 5.2, 5.3, 5.4 (stop at old 5.5)
    old52 = find_para(d, lambda p: p.text.strip().startswith("5.2  ชุดเครื่องมือโอเพนซอร์ส"))
    if old52:
        delete_block(old52, ["5.5"], d)

    # rename old 5.5 -> 5.3 and rewrite its body
    old55 = find_para(d, lambda p: p.text.strip().startswith("5.5"))
    if old55:
        set_text(old55, CH5_3_HEAD)
        body53 = docx.text.paragraph.Paragraph(old55._p.getnext(), d)
        set_text(body53, CH5_3_BODY)

    # insert new 5.2 (heading + body) after the roles paragraph
    if roles is not None and ref_h2 is not None:
        nb = new_para_after(roles._p, d, roles, CH5_2_BODY)
        nh = new_para_after(roles._p, d, ref_h2, CH5_2_HEAD)
        if nh._p.getnext() is not nb._p:
            nb._p.getparent().remove(nb._p)
            nh._p.addnext(nb._p)

    # -- global term change
    for p in d.paragraphs:
        if "เกตเวย์ขอบ" in p.text:
            set_text(p, p.text.replace("เกตเวย์ขอบ", "เอดจ์เกตเวย์"))

    d.save(DOC)
    print("Restructured chapters 3-5 (Thai); term เกตเวย์ขอบ -> เอดจ์เกตเวย์.")


if __name__ == '__main__':
    main()

#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Mirror the Thai chapter 3/4/5 restructure onto thesis_en.docx.

Chapter 3: rename 3.1/3.2, reframe 3.3-3.6 as software sections.
Chapter 4: rename 4.1 (drop "Study A:"), delete Study B (4.3) and table 4.2.
Chapter 5: rewrite into 5.1 System Overview / 5.2 Problems, Obstacles, and
           Solutions / 5.3 Recommendations (old 5.2-5.4 dropped, 5.5 -> 5.3).
Also: §2.6 "Monitoring" -> "Dashboard" stack (matches Thai ชุดแดชบอร์ด),
      objective 2 ingestion-comparison clause dropped, ch2 structure line and
      ch4 intro updated to drop the Node-RED vs Telegraf comparison.

§5.2 body is an AUTHOR-REVIEW DRAFT (marked inline).

Run once on the current thesis_en.docx:
    python3 build_scripts/restructure_ch345_en.py
"""
import copy
import os
import docx
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOC = os.path.join(ROOT, "thesis_en.docx")

RENAME = {
    "2.6  Time-Series Monitoring Stack (TIG)": "2.6  Time-Series Dashboard Stack (TIG)",
    "3.1  Overall Platform Architecture": "3.1  Overall Architecture",
    "3.2  Hardware Selection and Justification": "3.2  Hardware Selection: MCU, SBC, and PLC",
    "3.3  OT Layer: Siemens LOGO! PLC and LOGO! Soft Comfort":
        "3.3  OT-Layer Software: PLC Programming with LOGO! Soft Comfort",
    "3.4  Edge Gateway: ESP32-S3 Modbus Firmware":
        "3.4  Edge-Gateway Firmware: Modbus on ESP32-S3",
    "3.5  IT Layer: Containerised MQTT and TIG Stack":
        "3.5  IT-Layer Software: Containerised MQTT and TIG Stack",
    "4.1  Study A: Modbus Gateway Library Evaluation":
        "4.1  Modbus Gateway Library Evaluation",
}

REPLACE = [
    (", and comparing open-source ingestion pipelines for the IT layer", ""),
    (", gateway performance, and the Node-RED versus Telegraf comparison.",
     " and gateway performance."),
    (" the monitoring stack, and related work", " the dashboard stack, and related work"),
    (" and the monitoring stack. It closes", " and the dashboard stack. It closes"),
    ("connects the edge devices to the monitoring stack over MQTT",
     "connects the edge devices to the dashboard stack over MQTT"),
]

CH3_INTRO = ("This chapter presents the design and implementation of the platform. It "
             "begins with the overall architecture and the data paths, then explains the "
             "selection of the three hardware classes, and finally details the software "
             "of each layer: PLC programming, the edge-gateway firmware on the ESP32-S3, "
             "the containerised IT stack, and the IoT-Sniffer tool.")

CH4_INTRO = ("This chapter reports two evaluations: an empirical benchmark of open-source "
             "ESP32 Modbus gateway libraries, and the resource and timing performance of "
             "the resulting production gateway.")

CH5_INTRO = ("This chapter summarises the system that was built, discusses the problems "
             "and obstacles encountered during the work together with their solutions, "
             "and offers recommendations for further development.")

CH5_1_HEAD = "5.1  System Overview"
CH5_1_BODY = ("The platform integrates the three controller classes on a single "
              "local-area network. The ESP32-S3 microcontroller acts as the edge gateway, "
              "converting Modbus RTU to Modbus TCP and MQTT; the Raspberry Pi 4 "
              "single-board computer hosts the IT layer, running the MQTT broker and the "
              "TIG stack in containers; and the Siemens LOGO! 24CE PLC provides "
              "deterministic control. Data from the field instruments flows through the "
              "edge gateway up to the Grafana dashboards in real time. The results confirm "
              "that a single low-cost microcontroller can bridge legacy serial OT "
              "equipment to a modern IT dashboard within its firmware budget. Table 5.1 "
              "summarises the three controller classes side by side, showing that each "
              "device suits a different, complementary role rather than competing with the "
              "others.")

CH5_2_HEAD = "5.2  Problems, Obstacles, and Solutions"
CH5_2_BODY = ("(This section is a draft compiled from the project content; please review "
              "and adjust it to your actual experience.) Several problems and obstacles "
              "arose during design and development, together with their solutions. First, "
              "the open-source Modbus gateway libraries for the ESP32 are numerous and "
              "vary in quality, both in maintenance and in feature coverage, making "
              "selection difficult; this was addressed by defining clear evaluation "
              "criteria and benchmarking performance on real hardware to select "
              "systematically. Second, communication over the RS-485 bus is sensitive to "
              "wiring, A/B polarity, and termination, causing intermittent CRC errors; "
              "this was resolved by checking the wiring, sharing a common ground, and "
              "matching the baud rate across all devices. Third, the microcontroller's "
              "flash and RAM are limited, so combining several roles in one firmware "
              "required careful resource management; this was handled by choosing "
              "low-footprint libraries and continuously measuring memory use. Fourth, the "
              "multi-service IT layer (Mosquitto, InfluxDB, Telegraf, and Grafana) was "
              "complex and hard to reproduce; this was solved by declaring the whole "
              "configuration as code with Docker Compose for consistent redeployment. "
              "Fifth, hardware photographs from a phone camera were rotated by their EXIF "
              "data and displayed in the wrong orientation when imported; this was fixed "
              "by correcting the image orientation before insertion.")

CH5_3_HEAD = "5.3  Recommendations"
CH5_3_BODY = ("This work can be extended in several directions. On the edge gateway, a "
              "telemetry ring buffer should be added to tolerate Wi-Fi outages, and an "
              "over-the-air (OTA) firmware update mechanism should be added to ease "
              "maintenance. On the IT layer, scheduled InfluxDB backups, container health "
              "checks, and Grafana threshold alerting should be added to strengthen the "
              "reliability of the data plane. On the communication side, adding the "
              "OPC-UA protocol should be studied to broaden interoperability with other "
              "industrial systems. Finally, on adoption, the platform should be trialled "
              "with students in a laboratory course to assess its educational value, and "
              "taken further toward real deployment or higher-level research.")


def txt(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t'))).strip()


def set_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text); return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def find_para(doc, pred):
    for p in doc.paragraphs:
        sn = p.style.name.lower()
        if sn.startswith('toc') or sn == 'tableoffigures':
            continue
        if pred(p):
            return p
    return None


def new_para_after(anchor_el, doc, style_para, text):
    p_el = copy.deepcopy(style_para._p)
    for child in list(p_el):
        if child.tag == qn('w:r'):
            p_el.remove(child)
    anchor_el.addnext(p_el)
    para = docx.text.paragraph.Paragraph(p_el, doc)
    run = para.add_run(text)
    src = style_para.runs
    if src:
        rpr = src[0]._r.find(qn('w:rPr'))
        if rpr is not None:
            run._r.insert(0, copy.deepcopy(rpr))
    return para


def delete_block(start_para, stop_startswith, doc):
    el = start_para._p
    while el is not None:
        nxt = el.getnext()
        if el.tag == qn('w:p') and any(txt(el).startswith(s) for s in stop_startswith):
            break
        el.getparent().remove(el)
        el = nxt


def main():
    d = docx.Document(DOC)

    for p in d.paragraphs:
        t = p.text.strip()
        if t in RENAME:
            set_text(p, RENAME[t])

    for p in d.paragraphs:
        t = p.text
        new = t
        for old, rep in REPLACE:
            if old in new:
                new = new.replace(old, rep)
        if new != t:
            set_text(p, new)

    p = find_para(d, lambda p: p.text.strip().startswith("This chapter presents the platform from architecture"))
    if p: set_text(p, CH3_INTRO)
    p = find_para(d, lambda p: p.text.strip().startswith("This chapter reports three evaluations"))
    if p: set_text(p, CH4_INTRO)

    # chapter 4: delete Study B block (section + table 4.2)
    head43 = find_para(d, lambda p: p.text.strip().startswith("4.3  Study B"))
    if head43:
        delete_block(head43, ["Chapter 5"], d)

    # chapter 5 rewrite
    p = find_para(d, lambda p: p.text.strip().startswith("This chapter interprets the results"))
    if p: set_text(p, CH5_INTRO)
    h51 = find_para(d, lambda p: p.text.strip().startswith("5.1  MCU versus SBC"))
    if h51: set_text(h51, CH5_1_HEAD)
    b51 = find_para(d, lambda p: p.text.strip().startswith("The central design claim"))
    if b51: set_text(b51, CH5_1_BODY)
    ref_h2 = h51
    roles = find_para(d, lambda p: p.text.strip().startswith("The MCU delivers the lowest cost-per-channel"))

    old52 = find_para(d, lambda p: p.text.strip().startswith("5.2  Open-Source versus Proprietary Toolchains as a Teaching"))
    if old52:
        delete_block(old52, ["5.5"], d)

    old55 = find_para(d, lambda p: p.text.strip().startswith("5.5"))
    if old55:
        set_text(old55, CH5_3_HEAD)
        body53 = docx.text.paragraph.Paragraph(old55._p.getnext(), d)
        set_text(body53, CH5_3_BODY)

    if roles is not None and ref_h2 is not None:
        nb = new_para_after(roles._p, d, roles, CH5_2_BODY)
        nh = new_para_after(roles._p, d, ref_h2, CH5_2_HEAD)
        if nh._p.getnext() is not nb._p:
            nb._p.getparent().remove(nb._p)
            nh._p.addnext(nb._p)

    d.save(DOC)
    print("Restructured chapters 3-5 (English) to match the Thai edition.")


if __name__ == '__main__':
    main()

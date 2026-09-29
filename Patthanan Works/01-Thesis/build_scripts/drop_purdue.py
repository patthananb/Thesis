#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Remove all Purdue-model content from thesis_en.docx and thesis_th.docx.

Per author request: the Purdue reference model / level framework is dropped
entirely from both editions. IT/OT convergence discussion is kept; controllers
are described by role rather than by numbered Purdue level.

Removed / rewritten:
  - abstract, objective 1, chapter-2 and chapter-3 intros, related-work,
    chapter-5 discussion, and the conclusion: Purdue mentions stripped;
  - Section 2.1 heading renamed to "IT/OT Convergence" and the Purdue
    architecture paragraph deleted;
  - Thai figure 2.1 (the Purdue diagram) image + caption deleted;
  - the "Purdue level" row removed from the Chapter 5 comparison table.

Run as a late pipeline step, after rewrite_objectives.py / finalize_*:
    python3 build_scripts/drop_purdue.py
Substring-based, so re-running is safe.
"""
import os
import docx
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# ---- substring replacements (applied to any paragraph that contains the key)
REPLACE_TH = [
    ("ซึ่งผสานตัวควบคุมสามประเภทตามชั้นของแบบจำลอง Purdue ได้แก่",
     "ซึ่งผสานตัวควบคุมสามประเภท ได้แก่"),
    ("และตัวควบคุมทั้งสามประเภทสอดคล้องกับชั้น Purdue ที่แตกต่างกันสามารถทำงานร่วมกันได้",
     "และตัวควบคุมทั้งสามประเภทสามารถทำงานร่วมกันได้อย่างเสริมกัน"),
    ("(PLC, Siemens LOGO! 24CE) ตามชั้นของแบบจำลอง Purdue",
     "(PLC, Siemens LOGO! 24CE)"),
    ("ทบทวนการผสมผสานไอที/โอที แบบจำลอง Purdue ฮาร์ดแวร์",
     "ทบทวนการผสมผสานไอที/โอที ฮาร์ดแวร์"),
    ("ความแตกต่างระหว่างไอที/โอทีและแบบจำลองอ้างอิง Purdue จากนั้นสำรวจ",
     "ความแตกต่างระหว่างไอที/โอที จากนั้นสำรวจ"),
    ("2.1  การผสมผสานไอที/โอทีและแบบจำลองอ้างอิง Purdue",
     "2.1  การผสมผสานไอที/โอที"),
    ("และมีงานสำรวจโพรโทคอลไอไอโอทีและแบบจำลอง Purdue อยู่เป็นจำนวนมาก",
     "และมีงานสำรวจโพรโทคอลไอไอโอทีอยู่เป็นจำนวนมาก"),
    ("โดยเริ่มจากการกำหนดชั้น Purdue ให้กับอุปกรณ์ทุกตัวและนิยามเส้นทางข้อมูล",
     "โดยเริ่มจากการจัดวางอุปกรณ์ทุกตัวและนิยามเส้นทางข้อมูล"),
    ("อุปกรณ์แต่ละตัวเหมาะสมที่สุดในชั้น Purdue ที่ต่างกัน",
     "อุปกรณ์แต่ละตัวเหมาะสมที่สุดในบทบาทที่ต่างกัน"),
    ("การกำหนดชั้น Purdue จึงเป็นผลลัพธ์ทางวิศวกรรม",
     "การกำหนดบทบาทของอุปกรณ์แต่ละประเภทจึงเป็นผลลัพธ์ทางวิศวกรรม"),
    ("Siemens LOGO! 24CE ข้ามชั้น Purdue ระดับ 1 ถึง 3 บนเครือข่ายเดียว",
     "Siemens LOGO! 24CE บนเครือข่ายเดียว"),
]

REPLACE_EN = [
    ("map cleanly onto distinct Purdue layers as complementary, not competing, "
     "teaching subjects.",
     "serve complementary, not competing, roles as teaching subjects."),
    # the Thai edition's embedded English abstract is truncated mid-sentence at
    # "...distinct Purdue layers " — complete it and strip Purdue. Placed AFTER
    # the full-sentence key above so it cannot double-apply.
    ("the three controller classes map cleanly onto distinct Purdue layers ",
     "the three controller classes serve complementary, not competing, roles "
     "as teaching subjects."),
    ("(PLC, Siemens LOGO! 24CE) — across the layers of the Purdue model.",
     "(PLC, Siemens LOGO! 24CE)."),
    ("reviews IT/OT convergence, the Purdue model, the common hardware",
     "reviews IT/OT convergence, the common hardware"),
    ("It begins with the IT/OT divide and the Purdue reference model, surveys",
     "It begins with the IT/OT divide, surveys"),
    ("2.1  IT/OT Convergence and the Purdue Reference Model",
     "2.1  IT/OT Convergence"),
    ("surveys of IIoT protocols and the Purdue model are plentiful",
     "surveys of IIoT protocols are plentiful"),
    ("It first maps every device onto a Purdue level and defines the data "
     "paths, then justifies",
     "It first arranges the devices and defines the data paths, then justifies"),
    ("Each device is optimal at a different Purdue level for reasons",
     "Each device is optimal in a different role for reasons"),
    ("The Purdue assignment is therefore an engineering consequence, not a "
     "preference.",
     "The assignment of each device to its role is therefore an engineering "
     "consequence, not a preference."),
    ("a Siemens LOGO! 24CE PLC across Purdue Levels 1 to 3 on a single network",
     "a Siemens LOGO! 24CE PLC on a single network"),
]

# ---- whole-paragraph rewrites (match by unique opening substring)
REWRITE_TH = [
    ("แพลตฟอร์มบูรณาการตัวควบคุมทั้งสามประเภทบนเครือข่ายท้องถิ่นเดียว และจัดวางตามชั้น Purdue",
     "แพลตฟอร์มบูรณาการตัวควบคุมทั้งสามประเภทบนเครือข่ายท้องถิ่นเดียว เครื่องมือวัดหน้างานได้แก่ "
     "เซนเซอร์อุณหภูมิและความชื้น XY-MD02 และมิเตอร์ไฟฟ้าเฟสเดียว Eastron SDM230 ซึ่งทั้งสองเป็นสเลฟ "
     "Modbus RTU บนบัส RS-485 พีแอลซี Siemens LOGO! 24CE ทำหน้าที่ควบคุมที่มีความแน่นอนเชิงเวลาด้วย"
     "I/Oระดับอุตสาหกรรม เกตเวย์ขอบ ESP32-S3 ทำหน้าที่แปลงโพรโทคอล และ Raspberry Pi 4 รองรับ"
     "โบรกเกอร์ MQTT และชั้นข้อมูล TIG การออกแบบเปิดเผยเส้นทางข้อมูลสองเส้นทาง เส้นทาง A คือ Modbus TCP "
     "โดยตรง โฮสต์ใดก็ตามในเครือข่าย LAN รวมถึง PLC ที่ทำหน้าที่เป็นมาสเตอร์ สามารถสำรวจเซิร์ฟเวอร์ "
     "Modbus TCP ได้ เส้นทาง B คือไปป์ไลน์โทรมาตร ESP32-S3 สำรวจเซนเซอร์ RTU เผยแพร่ข้อมูล JSON ผ่าน "
     "MQTT และ Telegraf นำข้อมูลเข้าสู่ InfluxDB เพื่อแสดงผลใน Grafana ดังแสดงในภาพที่ 3.1"),
]

REWRITE_EN = [
    ("The platform integrates the three controller classes on a single "
     "local-area network and arranges them by Purdue level.",
     "The platform integrates the three controller classes on a single "
     "local-area network. The field instruments are an XY-MD02 temperature "
     "and humidity sensor and an Eastron SDM230 single-phase power meter, "
     "both Modbus RTU slaves on an RS-485 bus. The Siemens LOGO! 24CE PLC "
     "performs deterministic control with industrial I/O. The ESP32-S3 edge "
     "gateway converts protocols. The Raspberry Pi 4 hosts the MQTT broker "
     "and the TIG data plane. The design exposes two data paths. Path A is "
     "direct Modbus TCP: any LAN host, including the PLC acting as a master, "
     "polls a Modbus TCP server. Path B is the telemetry pipeline: the "
     "ESP32-S3 polls the RTU sensors, publishes JSON over MQTT, and Telegraf "
     "ingests it into InfluxDB for display in Grafana."),
]

# ---- whole paragraphs to delete (match by unique substring)
DELETE_TH = ["สถาปัตยกรรมอ้างอิงระดับองค์กร Purdue (Purdue Enterprise Reference Architecture)"]
DELETE_EN = ["The Purdue Enterprise Reference Architecture, adopted by the ISA-95 standard"]

# figure caption whose image+caption pair is deleted (Thai only)
DELETE_FIG_CAPTION_TH = "ภาพที่ 2.1  แบบจำลองอ้างอิงระดับองค์กร Purdue"


def set_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text); return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def remove_para(p):
    p._p.getparent().remove(p._p)


def process(fname, replaces, rewrites, deletes, fig_caption):
    path = os.path.join(ROOT, fname)
    d = docx.Document(path)

    # substring replacements
    for p in d.paragraphs:
        t = p.text
        new = t
        for old, rep in replaces:
            if old in new:
                new = new.replace(old, rep)
        if new != t:
            set_text(p, new)

    # whole-paragraph rewrites
    for p in d.paragraphs:
        for opening, full in rewrites:
            if p.text.strip().startswith(opening[:40]) and 'Purdue' in p.text:
                set_text(p, full)
                break

    # delete Purdue architecture paragraph
    for p in list(d.paragraphs):
        if any(key in p.text for key in deletes):
            remove_para(p)

    # delete figure image + caption (Thai) — must be the Caption-styled
    # paragraph, not the List-of-Figures cached field entry
    if fig_caption:
        for p in list(d.paragraphs):
            if p.style.name == 'Caption' and p.text.strip().startswith(fig_caption):
                prev = p._p.getprevious()
                if prev is not None and prev.tag == qn('w:p') and \
                        prev.find('.//' + qn('a:blip')) is not None:
                    prev.getparent().remove(prev)
                remove_para(p)
                break

    # delete the "Purdue level" row from the comparison table
    for t in d.tables:
        for row in list(t.rows):
            if 'Purdue' in row.cells[0].text:
                row._tr.getparent().remove(row._tr)

    d.save(path)
    left = sum(p.text.count('Purdue') for p in docx.Document(path).paragraphs)
    print(f"{fname}: done; 'Purdue' mentions left in body paragraphs = {left}")


if __name__ == '__main__':
    # the Thai edition also carries an English abstract page -> apply EN keys too
    process('thesis_th.docx', REPLACE_TH + REPLACE_EN, REWRITE_TH, DELETE_TH,
            DELETE_FIG_CAPTION_TH)
    process('thesis_en.docx', REPLACE_EN, REWRITE_EN, DELETE_EN, None)

#!/usr/bin/env python3
"""Correct the PLC from LOGO! 12/24RCE (relay) to LOGO! 24CE (transistor outputs)
across the combined report: Appendix E.2 specs + Table 3.2 and Table 4.3 cells.
Leaves the ESP32-S3 relay board untouched."""
from docx import Document

DOC = 'senior_project_report.docx'

REPL = [
    ("E.2  Siemens LOGO! 8.4 (6ED1052-1MD08-0BA2)",
     "E.2  Siemens LOGO! 24CE (firmware V8.4, 6ED1052-1CC08-0BA2)"),
    ("compact logic module (micro-PLC) in the 12/24 V DC / 24 V AC variant",
     "compact logic module (micro-PLC) in the 24 V DC variant (model 24CE, transistor outputs)"),
    ("6ED1052-1MD08-0BA2", "6ED1052-1CC08-0BA2"),
    ("Supply voltage: 12 VDC / 24 VDC / 24 VAC", "Supply voltage: 24 V DC"),
    ("4 × relay 10 A AC / 5 A DC (Q1–Q4)", "4 × transistor 24 V DC, 0.3 A (Q1–Q4)"),
    ("4x relay DO 10 A", "4x transistor DO 24 V DC 0.3 A"),
    ("relay DO 10 A (industrial levels)", "transistor DO 0.3 A (industrial levels)"),
    # Table 3.2 device-name cell
    ("Siemens LOGO! 8.4 (6ED1052-1CC08-0BA2)", "Siemens LOGO! 24CE (6ED1052-1CC08-0BA2)"),
]


def apply_to_paragraph(p):
    # try run-level first (handles the common single-run case, preserves formatting)
    for r in p.runs:
        if not r.text:
            continue
        for a, b in REPL:
            if a in r.text:
                r.text = r.text.replace(a, b)
    # fallback: target spans multiple runs -> rebuild paragraph text
    full = p.text
    if any(a in full for a, _ in REPL):
        new = full
        for a, b in REPL:
            new = new.replace(a, b)
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        p.add_run(new)


def main():
    doc = Document(DOC)
    for p in doc.paragraphs:
        apply_to_paragraph(p)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    apply_to_paragraph(p)
    doc.save(DOC)

    # verify
    d = Document(DOC)
    flat = '\n'.join(p.text for p in d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                flat += '\n' + c.text
    print('remaining MD08:', flat.count('6ED1052-1MD08-0BA2'))
    print('CC08 occurrences:', flat.count('6ED1052-1CC08-0BA2'))
    print('LOGO relay-10A left:', 'relay DO 10 A' in flat and 'LOGO' in flat, '| 4 × relay 10 A left:', '4 × relay 10 A' in flat)
    print('transistor mentions:', flat.count('transistor'))
    print('ESP32 SPDT relay intact:', '6x SPDT relay 10 A' in flat)


if __name__ == '__main__':
    main()

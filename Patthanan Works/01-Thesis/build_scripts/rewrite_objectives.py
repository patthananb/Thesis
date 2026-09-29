#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Realign thesis_en.docx and thesis_th.docx to four consolidated objectives.

Four objectives (per the author's summary):
  1. Create the IIoT laboratory platform (SBC + MCU + PLC).
  2. Demonstrate IT/OT integration with an open-source software stack.
  3. Evaluate ESP32-S3 Modbus gateway libraries (and measure the gateway).
  4. Develop the IoT-Sniffer to observe network activity and measure the
     network performance of the platform's IoT devices.

Objectives 4 (memory footprint) and 5 (Node-RED vs Telegraf) of the previous
six-item list are folded into objectives 3 and 2 respectively, so no chapter
is orphaned. The Section 5.4 conclusion and the Chapter 5 intro are rewritten
to match, and "six objectives" references become "four objectives".

Run after drop_security.py, before split_editions.py:
    python3 build_scripts/rewrite_objectives.py
"""
import os
import docx

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# ------------------------------------------------------------------ objectives
OBJ_EN = [
    "1)  Design and build a low-cost educational IIoT laboratory platform that "
    "integrates the three main hardware classes of modern automation — a "
    "single-board computer (SBC, Raspberry Pi 4), a microcontroller "
    "(MCU, ESP32-S3), and a programmable logic controller (PLC, Siemens LOGO! "
    "24CE) — across the layers of the Purdue model.",
    "2)  Demonstrate integration between IT and OT devices with an emphasis on "
    "an open-source software stack, adopting the tools most widely used in "
    "IIoT applications — the MQTT broker (Mosquitto), Node-RED, and the TIG "
    "stack (Telegraf, InfluxDB, Grafana) — as the platform's edge and IT "
    "layers, and comparing open-source ingestion pipelines for the IT layer.",
    "3)  Implement a Modbus RTU-to-TCP/MQTT gateway on the MCU and evaluate "
    "open-source ESP32-S3 Modbus gateway libraries, benchmarking the "
    "gateway-capable candidates on real hardware to select one, and measuring "
    "the resulting gateway's latency, throughput, and flash/RAM footprint.",
    "4)  Develop an IoT-Sniffer that decodes Modbus TCP and MQTT to observe "
    "network activity on the wire and measure the network performance of the "
    "platform's IoT devices.",
]

OBJ_TH = [
    "1)  ออกแบบและสร้างแพลตฟอร์มห้องปฏิบัติการไอโอทีเชิงอุตสาหกรรมต้นทุนต่ำเพื่อการศึกษา "
    "ที่บูรณาการฮาร์ดแวร์หลักสามประเภทของงานอัตโนมัติสมัยใหม่ ได้แก่ คอมพิวเตอร์บอร์ดเดี่ยว "
    "(SBC, Raspberry Pi 4) ไมโครคอนโทรลเลอร์ (MCU, ESP32-S3) และพีแอลซี "
    "(PLC, Siemens LOGO! 24CE) ตามชั้นของแบบจำลอง Purdue",
    "2)  สาธิตการบูรณาการระหว่างอุปกรณ์ไอทีและโอที โดยเน้นชุดซอฟต์แวร์โอเพนซอร์ส "
    "ด้วยการนำเครื่องมือที่ใช้กันแพร่หลายที่สุดในงานไอไอโอที ได้แก่ โบรกเกอร์ MQTT (Mosquitto), "
    "Node-RED และชุด TIG (Telegraf, InfluxDB, Grafana) มาใช้เป็นชั้นขอบและชั้นไอทีของแพลตฟอร์ม "
    "และเปรียบเทียบไปป์ไลน์รับข้อมูลโอเพนซอร์สสำหรับชั้นไอที",
    "3)  พัฒนาเกตเวย์ Modbus RTU เป็น TCP/MQTT บนไมโครคอนโทรลเลอร์ และประเมินไลบรารีเกตเวย์ "
    "Modbus โอเพนซอร์สบน ESP32-S3 โดยทดสอบสมรรถนะไลบรารีที่รองรับการทำงานเป็นเกตเวย์บนฮาร์ดแวร์จริง"
    "เพื่อคัดเลือกหนึ่งตัว และวัดเวลาแฝง อัตราการส่งข้อมูล และการใช้หน่วยความจำแฟลช/แรมของเกตเวย์ที่ได้",
    "4)  พัฒนาเครื่องมือ IoT-Sniffer ที่ถอดรหัส Modbus TCP และ MQTT เพื่อสังเกตกิจกรรมบนเครือข่าย"
    "และวัดสมรรถนะเครือข่ายของอุปกรณ์ไอโอทีในแพลตฟอร์ม",
]

# ------------------------------------------------------------------ conclusion
CONCL_EN = (
    "This work designed, built, and evaluated a low-cost educational Industrial "
    "IoT laboratory platform that integrates the three main hardware classes of "
    "modern automation. Against the four objectives: first, the platform "
    "integrates an ESP32-S3 microcontroller, a Raspberry Pi 4 single-board "
    "computer, and a Siemens LOGO! 24CE PLC across Purdue Levels 1 to 3 on a "
    "single network. Second, it demonstrates IT/OT integration on an "
    "open-source software stack — the MQTT broker, Node-RED, and the TIG stack — "
    "adopted as the platform's edge and IT layers, with Telegraf chosen over "
    "Node-RED for production ingestion. Third, a Modbus RTU-to-TCP/MQTT gateway "
    "was implemented on the microcontroller, and a benchmark of eight "
    "open-source ESP32 Modbus libraries selected emelianov/modbus-esp8266 on "
    "the strength of its maintenance, feature coverage, footprint, and "
    "robustness; the production gateway runs within 72.2 percent flash and 14.8 "
    "percent RAM on a single inexpensive board. Fourth, a purpose-built "
    "IoT-Sniffer decoded Modbus TCP and MQTT to observe network activity on the "
    "wire and measure the network performance of the platform's IoT devices, "
    "reporting throughput and p50/p95/p99 request latency. The platform thus "
    "turns the abstract goal of IT/OT convergence into a concrete, "
    "reproducible, and instructive laboratory system."
)

CONCL_TH = (
    "งานนี้ออกแบบ สร้าง และประเมินแพลตฟอร์มห้องปฏิบัติการไอโอทีเชิงอุตสาหกรรมต้นทุนต่ำเพื่อการศึกษา "
    "ที่บูรณาการฮาร์ดแวร์หลักสามประเภทของงานอัตโนมัติสมัยใหม่ เทียบกับวัตถุประสงค์ทั้งสี่ข้อ ข้อแรก "
    "แพลตฟอร์มบูรณาการไมโครคอนโทรลเลอร์ ESP32-S3 คอมพิวเตอร์บอร์ดเดี่ยว Raspberry Pi 4 และพีแอลซี "
    "Siemens LOGO! 24CE ข้ามชั้น Purdue ระดับ 1 ถึง 3 บนเครือข่ายเดียว ข้อที่สอง แพลตฟอร์มสาธิตการ"
    "บูรณาการไอที/โอทีบนชุดซอฟต์แวร์โอเพนซอร์ส ได้แก่ โบรกเกอร์ MQTT, Node-RED และชุด TIG ซึ่งนำมาใช้"
    "เป็นชั้นขอบและชั้นไอทีของแพลตฟอร์ม โดยเลือก Telegraf แทน Node-RED สำหรับการนำเข้าข้อมูลเชิงผลิต "
    "ข้อที่สาม ได้พัฒนาเกตเวย์ Modbus RTU เป็น TCP/MQTT บนไมโครคอนโทรลเลอร์ และการทดสอบไลบรารี Modbus "
    "โอเพนซอร์สสำหรับ ESP32 แปดตัวคัดเลือก emelianov/modbus-esp8266 ด้วยจุดแข็งด้านการบำรุงรักษา "
    "ความครอบคลุมคุณสมบัติ ขนาดการใช้ทรัพยากร และความทนทาน โดยเกตเวย์เชิงผลิตทำงานภายในแฟลชร้อยละ 72.2 "
    "และแรมร้อยละ 14.8 บนบอร์ดราคาประหยัดเพียงบอร์ดเดียว ข้อที่สี่ เครื่องมือ IoT-Sniffer ที่สร้างขึ้น"
    "เฉพาะถอดรหัส Modbus TCP และ MQTT เพื่อสังเกตกิจกรรมบนเครือข่ายและวัดสมรรถนะเครือข่ายของอุปกรณ์"
    "ไอโอทีในแพลตฟอร์ม โดยรายงานอัตราการส่งข้อมูลและเวลาแฝงของคำขอที่ p50/p95/p99 แพลตฟอร์มนี้จึงเปลี่ยน"
    "เป้าหมายเชิงนามธรรมของการหลอมรวมไอที/โอทีให้กลายเป็นระบบห้องปฏิบัติการที่เป็นรูปธรรม ทำซ้ำได้ และให้ความรู้"
)

# ch5 intro sentence
INTRO_FIND = [
    ("concludes against the six objectives", "concludes against the four objectives"),
    ("สรุปผลเทียบกับวัตถุประสงค์ทั้งหกข้อ", "สรุปผลเทียบกับวัตถุประสงค์ทั้งสี่ข้อ"),
]

EDITIONS = {
    'thesis_en.docx': (OBJ_EN, CONCL_EN, "5.4  Conclusion"),
    'thesis_th.docx': (OBJ_TH, CONCL_TH, "5.4  บทสรุป"),
}


def set_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def process(fname, objectives, conclusion, concl_heading):
    path = os.path.join(ROOT, fname)
    d = docx.Document(path)
    paras = d.paragraphs

    # locate the six objective list items (paragraphs starting "1) ".."6) ")
    obj_idx = {}
    for i, p in enumerate(paras):
        t = p.text.lstrip()
        for n in range(1, 7):
            if t.startswith(f"{n})  ") or t.startswith(f"{n}) "):
                obj_idx.setdefault(n, i)
    # rewrite 1..4, delete 5 and 6
    for n in range(1, 5):
        set_text(paras[obj_idx[n]], objectives[n - 1])
    for n in (5, 6):
        if n in obj_idx:
            p = paras[obj_idx[n]]
            p._p.getparent().remove(p._p)

    # ch5 intro count fix
    for p in d.paragraphs:
        for old, new in INTRO_FIND:
            if old in p.text:
                set_text(p, p.text.replace(old, new))

    # rewrite the conclusion paragraph (first Normal para after the 5.4 heading)
    paras = d.paragraphs
    hi = next(i for i, p in enumerate(paras) if p.text.strip() == concl_heading)
    for j in range(hi + 1, len(paras)):
        if paras[j].text.strip():
            set_text(paras[j], conclusion)
            break

    d.save(path)
    print(f"{fname}: objectives -> 4, conclusion rewritten.")


if __name__ == '__main__':
    for fname, (obj, concl, head) in EDITIONS.items():
        if os.path.exists(os.path.join(ROOT, fname)):
            process(fname, obj, concl, head)

#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Finalize the English edition front matter.

1. Make body page numbering continuous (chapters no longer restart at 1),
   matching the Thai edition, so the List of Tables page numbers are unique.
2. Populate the (empty) List of Tables with a real Word field (TOC \\f T) plus
   TC markers on the four table captions, AND a cached, immediately-visible
   result with dot-leader page numbers. Open in Word and press Ctrl+A, F9 to
   recompute the numbers if the layout later shifts.

Run after rewrite_objectives.py, before split_editions.py:
    python3 build_scripts/finalize_thesis_en.py
"""
import os
import docx
from docx.shared import Pt, Twips
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOC = os.path.join(ROOT, "thesis_en.docx")

# List-of-Tables entries: (caption text, printed page number).
# Page numbers are the cached display values; F9 recomputes them in Word.
ENTRIES = [
    ("Table 2.1  Representative technology stacks in Information Technology "
     "and Operational Technology", "6"),
    ("Table 4.1  FC03 round-trip latency and resource usage per Modbus "
     "library (ESP32-S3, 100 registers, 9,600 baud)", "17"),
    ("Table 4.2  Telegraf versus Node-RED feature comparison", "19"),
    ("Table 5.1  Comparison of the three controller classes used in the "
     "platform", "21"),
]
RIGHT_TAB_TWIPS = 8300  # ~ right margin for A4 with 1.5in/1in margins


def add_fldchar(run, kind, dirty=False):
    fc = OxmlElement('w:fldChar')
    fc.set(qn('w:fldCharType'), kind)
    if dirty:
        fc.set(qn('w:dirty'), 'true')
    run._element.append(fc)


def add_instr(run, text):
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = text
    run._element.append(it)


def style_entry(p):
    p.style = 'TOC0' if 'TOC0' in [s.name for s in p.part.document.styles] else 'Normal'
    pf = p.paragraph_format
    pf.left_indent = Twips(0)
    pf.first_line_indent = Twips(0)
    p.paragraph_format.tab_stops.add_tab_stop(
        Twips(RIGHT_TAB_TWIPS), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)


def new_para_after(anchor_el, doc):
    p = OxmlElement('w:p')
    anchor_el.addnext(p)
    return docx.text.paragraph.Paragraph(p, doc)


def main():
    d = docx.Document(DOC)

    # 1. continuous body numbering (sections 3..6 = chapters 2..5)
    secs = d.sections
    for si in (3, 4, 5, 6):
        pg = secs[si]._sectPr.find(qn('w:pgNumType'))
        if pg is not None and pg.get(qn('w:start')) is not None:
            del pg.attrib[qn('w:start')]

    # 2. locate the empty List-of-Tables field shell (Heading0 para right after
    #    the "List of Tables" heading) and the four captions.
    paras = d.paragraphs
    lot_head = next(i for i, p in enumerate(paras)
                    if p.text.strip() == 'List of Tables')
    # the shell paragraph: first paragraph after the heading that carries a
    # fldChar or is otherwise empty before "Nomenclature"
    nomen = next(i for i, p in enumerate(paras) if p.text.strip() == 'Nomenclature')
    shell = None
    for i in range(lot_head + 1, nomen):
        if paras[i]._p.findall('.//' + qn('w:fldChar')):
            shell = paras[i]
            break
    if shell is None:
        shell = paras[lot_head + 1]
    # clear the shell paragraph contents and un-Heading it
    for r in list(shell.runs):
        r._r.getparent().remove(r._r)
    for child in list(shell._p):
        if child.tag == qn('w:r'):
            shell._p.remove(child)
    shell.style = 'Normal'

    # build the field: begin + instr + separate in the shell, then one
    # paragraph per entry, end fldChar in the last entry paragraph.
    r = shell.add_run(); add_fldchar(r, 'begin', dirty=True)
    r = shell.add_run(); add_instr(r, ' TOC \\h \\z \\f T ')
    r = shell.add_run(); add_fldchar(r, 'separate')

    anchor = shell._p
    for k, (text, page) in enumerate(ENTRIES):
        p = new_para_after(anchor, d)
        style_entry(p)
        p.add_run(text)
        p.add_run('\t')
        pr = p.add_run(page)
        if k == len(ENTRIES) - 1:
            end = p.add_run(); add_fldchar(end, 'end')
        anchor = p._p

    # 3. TC markers on the four captions so a future F9 rebuilds the list
    for p in d.paragraphs:
        t = p.text.strip()
        if p.style.name == 'Caption' and t.startswith('Table'):
            r1 = p.add_run(); add_fldchar(r1, 'begin')
            r2 = p.add_run(); add_instr(r2, ' TC "%s" \\f T \\l "1" ' %
                                        t.rstrip('.').replace('"', "'"))
            r3 = p.add_run(); add_fldchar(r3, 'end')

    d.save(DOC)
    print("English List of Tables populated; body numbering continuous.")


if __name__ == '__main__':
    main()

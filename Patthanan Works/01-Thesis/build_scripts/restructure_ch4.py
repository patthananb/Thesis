#!/usr/bin/env python3
"""Restructure Chapter 4 of the English report:
 - New Study A = IIoT Platform Comparison (move the Appendix F matrix into Ch4 as Table 4.1)
 - Study B = Modbus library evaluation (was Study A); move its flash/RAM table into it
 - Study C = Telegraf vs Node-RED (was Study B); replace content with the new comparison
 - Security -> 4.4, MCU/SBC/PLC -> 4.5; tables/sections/refs renumbered; Appendix F removed.
"""
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor

DOC = 'senior_project_report.docx'

# ---- global text remaps (sequential, descending so each token maps once) ----
REMAPS = [
    ('Section 4.4', 'Section 4.5'), ('Section 4.3', 'Section 4.4'),
    ('Section 4.2', 'Section 4.3'), ('Section 4.1', 'Section 4.2'),
    ('Table 4.3', 'Table 4.5'), ('Table 4.2', 'Table 4.3'), ('Table 4.1', 'Table 4.2'),
    ('Study B', 'Study C'), ('Study A', 'Study B'),
]

HEAD = {  # final heading text by original prefix
    '4.1  Study A: Modbus': '4.2  Study B: Modbus TCP Gateway Library Evaluation',
    '4.1.1': '4.2.1  Setup and Criteria',
    '4.1.2': '4.2.2  Results',
    '4.2  Study B: Node-RED': '4.3  Study C: Telegraf vs Node-RED',
    '4.2.1': '4.3.1  Core Philosophy and Performance',
    '4.2.2': '4.3.2  Application Fit and Recommendation',
    '4.3  Security': '4.4  Security Demonstration',
    '4.4  MCU': '4.5  MCU vs SBC vs PLC: Role Comparison',
    '4.4.1': '4.5.1  Microcontroller Unit (MCU) — ESP32-S3',
    '4.4.2': '4.5.2  Single-Board Computer (SBC) — Raspberry Pi 4',
    '4.4.3': '4.5.3  Programmable Logic Controller (PLC) — Siemens LOGO! 8.4',
    '4.4.4': '4.5.4  Summary: Right Device at the Right Purdue Level',
}

STUDYA_FIND = (
    "The matrix confirms the device roles adopted in this project: the ESP32-S3 (MCU) "
    "excels at field-edge I/O and protocol bridging at the lowest cost and power; the "
    "Raspberry Pi 4 (SBC) is the only platform offering the full Linux and Docker stack "
    "for the data plane; and the Siemens LOGO! 24CE (PLC) is the only deterministic, "
    "industrially certified option for control. No single device wins every row, which "
    "is precisely what motivates the layered IT/OT architecture used here.")

STUDYC_P1 = (
    "Both Telegraf and Node-RED can ingest data (via MQTT or HTTP) and write it to a "
    "database such as InfluxDB, yet they fill fundamentally different architectural roles. "
    "This study compares them as the MQTT-to-InfluxDB pipeline for the project.")
STUDYC_PHIL = (
    "Telegraf is a server agent written in Go — a lightweight, set-and-forget compiled "
    "binary configured through a single .conf file and executed at scale. Node-RED is a "
    "visual programming environment on Node.js that emphasises the flow of data, letting "
    "the developer drag and drop nodes to build event-driven logic.")
STUDYC_PERF = (
    "In performance, Telegraf has very low overhead and a tiny memory footprint, sustains "
    "high throughput (thousands of metrics per second), and buffers internally so that "
    "data is cached in RAM if InfluxDB is briefly unavailable. Node-RED, as a Node.js "
    "application, uses more RAM and CPU, and the overhead of its visual engine makes it "
    "less suited to firehose-style ingestion, though it is fast enough for typical IoT "
    "workloads.")
STUDYC_FIT = (
    "Telegraf suits infrastructure monitoring (CPU, disk, and network across many hosts), "
    "standardised IIoT pass-through from OPC-UA, Modbus, or MQTT into a database, and edge "
    "computing on resource-constrained devices where every megabyte of RAM counts. "
    "Node-RED suits home and building automation logic (“if sensor A > 25 °C and the "
    "time is after 8 PM, switch on the AC and send a Telegram message”), rapid "
    "prototyping of an API or a new sensor, and bridging legacy protocols with custom "
    "parsing in between.")
STUDYC_SYN = (
    "The two tools are complementary. In advanced stacks Telegraf does the heavy lifting — "
    "ingesting high-resolution telemetry and system metrics into InfluxDB — while Node-RED "
    "acts as the brain, querying InfluxDB or specific MQTT topics to trigger alerts, "
    "control actuators, and manage dashboards. For this project's production pipeline, "
    "Telegraf was selected for ingestion (low idle memory of roughly 20–40 MB versus "
    "80–120 MB, internal buffering, and a Git-friendly TOML configuration), with "
    "Node-RED retained for event logic and rapid prototyping.")

FEATURE_ROWS = [
    ['Feature', 'Telegraf', 'Node-RED'],
    ['Interface', 'CLI / configuration file (TOML)', 'Browser-based visual UI'],
    ['Logic', 'Fixed: collect → process → aggregate → output',
     'Freeform: conditional branching, loops, JavaScript functions'],
    ['Customization', 'Write Go plugins (recompilation)', 'Drag in new nodes or inline JavaScript'],
    ['Deployment', 'Single binary, easy to automate (Ansible/Docker)',
     'Requires Node.js; flows harder to version-control'],
]


def starts(p, pre):
    return p.text.strip().startswith(pre)


def set_text(p, text, bold=False):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(text)
    run.bold = bold


def move_after(anchor_el, elems):
    ref = anchor_el
    for e in elems:
        e.getparent().remove(e)
        ref.addnext(e)
        ref = e
    return ref


def move_before(anchor_el, elems):
    for e in elems:
        e.getparent().remove(e)
        anchor_el.addprevious(e)


def find(doc, pred):
    # body-only: skip TOC / List-of-Figures / List-of-Tables entries (they carry a tab)
    for p in doc.paragraphs:
        if p._p.find('.//' + qn('w:tab')) is not None:
            continue
        if pred(p):
            return p
    return None


def main():
    doc = Document(DOC)

    # ---- capture references BEFORE editing ----
    P = doc.paragraphs
    ch4 = find(doc, lambda p: starts(p, 'CHAPTER 4'))
    modbus_head = find(doc, lambda p: starts(p, '4.1  Study A: Modbus'))
    nodered_head = find(doc, lambda p: starts(p, '4.2  Study B: Node-RED'))
    h421 = find(doc, lambda p: starts(p, '4.2.1'))
    h422 = find(doc, lambda p: starts(p, '4.2.2'))
    setup_para = find(doc, lambda p: starts(p, 'Both pipelines subscribe'))
    telegraf_sum = find(doc, lambda p: starts(p, 'Telegraf outperforms Node-RED'))
    flash_cap = find(doc, lambda p: starts(p, 'Table 4.2  Study A: Flash'))
    flash_d1 = find(doc, lambda p: starts(p, 'Table 4.2 shows flash'))
    flash_d2 = find(doc, lambda p: starts(p, 'The three raw-RTU'))
    flash_d3 = find(doc, lambda p: starts(p, 'For the production gateway'))
    appf_head = find(doc, lambda p: starts(p, 'APPENDIX F'))
    appf_intro = find(doc, lambda p: starts(p, 'Table F.1 compares'))
    appf_cap = find(doc, lambda p: starts(p, 'Table F.1  IIoT platform'))
    appf_legend = find(doc, lambda p: p.text.strip().startswith('★') and 'strong' in p.text)
    heads = {k: find(doc, lambda p, k=k: starts(p, k)) for k in HEAD}

    flash_table = flash_cap._p.getprevious()        # table sits above the caption
    appf_table = appf_cap._p.getnext()              # table sits below the caption

    # ---- 1) global text remaps (skip TOC/list tab rows) ----
    targets = list(doc.paragraphs)
    for tb in doc.tables:
        for row in tb.rows:
            for c in row.cells:
                targets.extend(c.paragraphs)
    for p in targets:
        if p._p.find('.//' + qn('w:tab')) is not None:
            continue
        for r in p.runs:
            if not r.text:
                continue
            for a, b in REMAPS:
                if a in r.text:
                    r.text = r.text.replace(a, b)

    # ---- 2) move flash/RAM block into Modbus (Study B), just before Node-RED heading ----
    move_before(nodered_head._p, [flash_cap._p, flash_table, flash_d1._p, flash_d2._p, flash_d3._p])

    # ---- 3) delete Node-RED leftover content (setup para + telegraf summary) ----
    for el in (setup_para._p, telegraf_sum._p):
        el.getparent().remove(el)

    # ---- 4) rename headings to final text ----
    for k, p in heads.items():
        set_text(p, HEAD[k], bold=True)

    # ---- 5) Study C new content ----
    # after 4.3.1 heading -> philosophy + performance
    a = move_after(h421._p, [doc.add_paragraph(STUDYC_P1)._p,
                             doc.add_paragraph(STUDYC_PHIL)._p,
                             doc.add_paragraph(STUDYC_PERF)._p])
    # after 4.3.2 heading -> caption + feature table + fit + synergy
    cap = doc.add_paragraph('Table 4.4  Telegraf vs Node-RED: feature comparison.')
    tbl = doc.add_table(rows=len(FEATURE_ROWS), cols=3)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(FEATURE_ROWS):
        for ci, val in enumerate(row):
            cell = tbl.rows[ri].cells[ci]
            cell.text = ''
            run = cell.paragraphs[0].add_run(val)
            if ri == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                tcPr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:fill'), '0F2A43'); tcPr.append(shd)
    fit = doc.add_paragraph(STUDYC_FIT)
    syn = doc.add_paragraph(STUDYC_SYN)
    move_after(h422._p, [cap._p, tbl._tbl, fit._p, syn._p])

    # ---- 6) build Study A from the Appendix F matrix ----
    set_text(appf_cap, 'Table 4.1  IIoT platform comparison matrix (MCU vs SBC vs PLC).')
    for r in appf_intro.runs:
        if r.text and 'Table F.1' in r.text:
            r.text = r.text.replace('Table F.1', 'Table 4.1')
    sa_head = doc.add_paragraph(); set_text(sa_head, '4.1  Study A: IIoT Platform Comparison', bold=True)
    sa_find = doc.add_paragraph(STUDYA_FIND)
    move_after(ch4._p, [sa_head._p, appf_intro._p, appf_cap._p, appf_table, appf_legend._p, sa_find._p])

    # ---- 7) remove the Appendix F heading (content already relocated) ----
    appf_head._p.getparent().remove(appf_head._p)

    doc.save(DOC)

    # ---- verify ----
    d = Document(DOC)
    print('Ch4 outline:')
    inch = False
    for p in d.paragraphs:
        if '\t' in p.text:
            continue
        t = p.text.strip()
        if t.startswith('CHAPTER 4'):
            inch = True
        if t.startswith('CHAPTER 5'):
            inch = False
        if inch and (t[:3] in ('4.1', '4.2', '4.3', '4.4', '4.5') or t.startswith('Table 4')):
            print('  ', t[:70])
    print('Appendix F remains:', any(p.text.strip().startswith('APPENDIX F') for p in d.paragraphs))
    print('tables:', len(d.tables))


if __name__ == '__main__':
    main()

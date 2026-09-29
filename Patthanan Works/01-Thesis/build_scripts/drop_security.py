#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Remove the security content from thesis_en.docx and thesis_th.docx.

The IoT-Sniffer is KEPT, reframed as a wire-level observation / teaching tool.
Removed: §2.7 (security challenges), §4.4 (security demonstration), the
security half of objective 6, the third research gap, the security mentions in
abstract/scope/structure/discussion/future work, and references [9]-[10].
§2.8 is renumbered to §2.7.

Run as the LAST build step, after finalize_thesis_th.py:
    python3 build_scripts/drop_security.py
(Text-matching, so it is safe to run on either the fresh or the finalized
documents; anything already clean is left untouched.)
"""
import os
import docx
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# --------------------------------------------------------------- edit tables
# headings whose whole section (heading + body until next heading) is removed
DELETE_SECTIONS = [
    "2.7  Security Challenges in IT/OT Environments",
    "2.7  ความท้าทายด้านความปลอดภัยในสภาพแวดล้อมไอที/โอที",
    "4.4  Security Demonstration",
    "4.4  การสาธิตความปลอดภัย",
]

# substring replacements, applied to any paragraph that contains the key
REPLACE = [
    # ---- abstract ----
    ("An in-house passive sniffer decodes Modbus TCP and MQTT on the wire "
     "and is used to demonstrate the unauthenticated-actuation and "
     "anonymous-publish gaps of these protocols.",
     "An in-house passive sniffer decodes Modbus TCP and MQTT on the wire, "
     "giving learners direct visibility into live protocol behaviour."),
    ("และได้พัฒนาเครื่องมือดักจับแพ็กเก็ตเชิงรับสำหรับสาธิตช่องโหว่การสั่งงานโดยไม่ต้องยืนยันตัวตนของ "
     "Modbus TCP และ MQTT",
     "และได้พัฒนาเครื่องมือดักจับแพ็กเก็ตเชิงรับสำหรับสังเกตพฤติกรรมของ Modbus TCP และ MQTT "
     "บนสายสัญญาณแบบสด"),
    # ---- 1.1 background ----
    ("with no authentication, no encryption, and no native path to an IP network",
     "with no native path to an IP network"),
    ("โดยไม่มีการยืนยันตัวตน ไม่มีการเข้ารหัส และไม่มีเส้นทางเชื่อมต่อเครือข่ายไอพีโดยกำเนิด",
     "โดยไม่มีเส้นทางเชื่อมต่อเครือข่ายไอพีโดยกำเนิด"),
    # ---- 1.2 research gaps: three -> two ----
    ("Three gaps follow.", "Two gaps follow."),
    (" Third, the insecurity of Modbus and MQTT is widely asserted but seldom "
     "shown on real traffic in a classroom-safe way.", ""),
    ("จึงเกิดช่องว่างสามประการ", "จึงเกิดช่องว่างสองประการ"),
    (" ประการที่สาม ความไม่ปลอดภัยของ Modbus และ MQTT ถูกกล่าวอ้างอย่างกว้างขวางแต่ไม่ค่อยถูกสาธิต"
     "บนการสื่อสารจริงในรูปแบบที่ปลอดภัยสำหรับห้องเรียน", ""),
    # ---- objective 6 ----
    ("and use it to demonstrate the IT/OT security gaps of these protocols.",
     "and use it for wire-level observation of the platform's traffic."),
    ("และใช้เครื่องมือนี้สาธิตช่องโหว่ด้านความปลอดภัยของไอที/โอทีในโพรโทคอลทั้งสอง",
     "และใช้เครื่องมือนี้สังเกตพฤติกรรมการสื่อสารระดับสายสัญญาณของแพลตฟอร์ม"),
    # ---- 1.4 scope ----
    ("the TIG stack, the passive sniffer, and a live security demonstration.",
     "the TIG stack, and the passive sniffer."),
    (" Security mitigations are discussed and partially implemented but not "
     "exhaustively deployed.", ""),
    ("ชุด TIG, เครื่องมือดักจับเชิงรับ และการสาธิตความปลอดภัยแบบสด",
     "ชุด TIG และเครื่องมือดักจับเชิงรับ"),
    (" มาตรการลดความเสี่ยงด้านความปลอดภัยถูกอภิปรายและนำไปใช้เพียงบางส่วน แต่ไม่ได้ติดตั้งอย่างครบถ้วน",
     ""),
    # ---- 1.5 thesis structure ----
    ("the Node-RED versus Telegraf comparison, and the security demonstration.",
     "and the Node-RED versus Telegraf comparison."),
    ("การเปรียบเทียบ Node-RED กับ Telegraf และการสาธิตความปลอดภัย บทที่ 5",
     "และการเปรียบเทียบ Node-RED กับ Telegraf บทที่ 5"),
    # ---- ch2 intro ----
    ("the monitoring stack, and the security properties of the protocols "
     "involved. It closes",
     "and the monitoring stack. It closes"),
    ("ชุดเฝ้าระวัง และคุณสมบัติด้านความปลอดภัยของโพรโทคอลที่เกี่ยวข้อง ปิดท้ายด้วย",
     "และชุดเฝ้าระวัง ปิดท้ายด้วย"),
    # ---- 2.3.1 Modbus RTU ----
    ("Modbus RTU has no authentication, no encryption, and no addressing "
     "beyond the slave id, which makes it simple to implement but trivially "
     "open on the wire.",
     "Modbus RTU has no addressing beyond the slave id, which keeps the "
     "protocol simple to implement."),
    ("Modbus RTU ไม่มีการยืนยันตัวตน ไม่มีการเข้ารหัส และไม่มีการระบุที่อยู่นอกเหนือจากหมายเลขสเลฟ "
     "ทำให้ง่ายต่อการนำไปใช้งานแต่เปิดเผยข้อมูลบนสายสัญญาณอย่างง่ายดาย",
     "Modbus RTU ไม่มีการระบุที่อยู่นอกเหนือจากหมายเลขสเลฟ ทำให้โพรโทคอลเรียบง่ายและนำไปใช้งานได้ง่าย"),
    # ---- 2.3.3 MQTT ----
    ("MQTT brokers permit anonymous access by default, a convenience that, "
     "as Section 2.7 discusses, is also a hazard.",
     "MQTT brokers permit anonymous access by default, which simplifies "
     "classroom deployment."),
    ("โบรกเกอร์ MQTT อนุญาตให้เข้าถึงแบบไม่ระบุตัวตนโดยค่าเริ่มต้น ซึ่งเป็นความสะดวกสบายที่ในขณะ"
     "เดียวกันก็เป็นความเสี่ยง ดังที่กล่าวถึงในหัวข้อ 2.7",
     "โบรกเกอร์ MQTT อนุญาตให้เข้าถึงแบบไม่ระบุตัวตนโดยค่าเริ่มต้น ซึ่งช่วยให้การติดตั้งใช้งานใน"
     "ห้องเรียนทำได้ง่าย"),
    # ---- 2.8 related work: three gaps -> two ----
    ("However, three gaps remain.", "However, two gaps remain."),
    (" Third, the insecurity of Modbus and MQTT is asserted in the literature "
     "but rarely shown in a reproducible, classroom-safe demonstration.", ""),
    ("This thesis addresses all three: it integrates the three controller "
     "classes on one low-cost bench, benchmarks eight of ten identified "
     "Modbus libraries to select a gateway implementation, and demonstrates "
     "the protocol weaknesses live with a purpose-built sniffer.",
     "This thesis addresses both: it integrates the three controller classes "
     "on one low-cost bench, and benchmarks eight of ten identified Modbus "
     "libraries to select a gateway implementation."),
    ("ยังมีช่องว่างอยู่สามประการ", "ยังมีช่องว่างอยู่สองประการ"),
    (" ประการที่สาม ความไม่ปลอดภัยของ Modbus และ MQTT ถูกกล่าวอ้างในวรรณกรรมแต่ไม่ค่อยถูกแสดง"
     "ให้เห็นในรูปแบบที่ทำซ้ำได้และปลอดภัยสำหรับห้องเรียน", ""),
    ("ปริญญานิพนธ์นี้แก้ไขช่องว่างทั้งสามประการ โดยบูรณาการตัวควบคุมทั้งสามประเภทบนชุดทดลองต้นทุนต่ำ"
     "ชุดเดียว ทดสอบสมรรถนะไลบรารี Modbus แปดในสิบตัวที่ระบุไว้เพื่อคัดเลือกการนำไปใช้เป็นเกตเวย์ "
     "และสาธิตจุดอ่อนของโพรโทคอลแบบสดด้วยเครื่องมือดักจับที่สร้างขึ้นเฉพาะ",
     "ปริญญานิพนธ์นี้แก้ไขช่องว่างทั้งสองประการ โดยบูรณาการตัวควบคุมทั้งสามประเภทบนชุดทดลองต้นทุนต่ำ"
     "ชุดเดียว และทดสอบสมรรถนะไลบรารี Modbus แปดในสิบตัวที่ระบุไว้เพื่อคัดเลือกการนำไปใช้เป็นเกตเวย์"),
    # ---- 3.6 sniffer purpose ----
    ("The tool serves two purposes: during development it verifies byte order "
     "and protocol behaviour on the wire, and during the security "
     "demonstration of Chapter 4 it exposes unauthenticated Modbus writes and "
     "anonymous MQTT publishes as plain, readable frames.",
     "The tool verifies byte order and protocol behaviour on the wire during "
     "development, and gives learners a live, readable view of every Modbus "
     "and MQTT frame the platform exchanges."),
    ("เครื่องมือนี้มีสองวัตถุประสงค์ ในระหว่างการพัฒนามันตรวจสอบลำดับไบต์และพฤติกรรมโพรโทคอลบนสาย"
     "สัญญาณ และในระหว่างการสาธิตความปลอดภัยในบทที่ 4 มันเปิดเผยการเขียน Modbus โดยไม่ยืนยันตัวตน"
     "และการเผยแพร่ MQTT แบบไม่ระบุตัวตนให้เห็นเป็นเฟรมที่อ่านได้ชัดเจน",
     "เครื่องมือนี้ใช้ตรวจสอบลำดับไบต์และพฤติกรรมโพรโทคอลบนสายสัญญาณระหว่างการพัฒนา และช่วยให้"
     "ผู้เรียนเห็นทุกเฟรม Modbus และ MQTT ที่แพลตฟอร์มแลกเปลี่ยนกันแบบสดในรูปแบบที่อ่านได้ชัดเจน"),
    # ---- ch4 intro: four -> three evaluations ----
    ("This chapter reports four evaluations: an empirical benchmark of "
     "open-source ESP32 Modbus libraries (Study A), the resource and timing "
     "performance of the resulting production gateway, a comparison of two "
     "open-source ingestion pipelines (Study B), and a live demonstration of "
     "the security gaps of the platform's protocols.",
     "This chapter reports three evaluations: an empirical benchmark of "
     "open-source ESP32 Modbus libraries (Study A), the resource and timing "
     "performance of the resulting production gateway, and a comparison of "
     "two open-source ingestion pipelines (Study B)."),
    ("บทนี้รายงานการประเมินสี่ส่วน ได้แก่ การทดสอบเชิงประจักษ์ของไลบรารี Modbus โอเพนซอร์สสำหรับ "
     "ESP32 (การศึกษา A) สมรรถนะด้านทรัพยากรและเวลาของเกตเวย์เชิงผลิตที่ได้ การเปรียบเทียบไปป์ไลน์"
     "รับข้อมูลโอเพนซอร์สสองแบบ (การศึกษา B) และการสาธิตสดของช่องโหว่ด้านความปลอดภัยในโพรโทคอล"
     "ของแพลตฟอร์ม",
     "บทนี้รายงานการประเมินสามส่วน ได้แก่ การทดสอบเชิงประจักษ์ของไลบรารี Modbus โอเพนซอร์สสำหรับ "
     "ESP32 (การศึกษา A) สมรรถนะด้านทรัพยากรและเวลาของเกตเวย์เชิงผลิตที่ได้ และการเปรียบเทียบ"
     "ไปป์ไลน์รับข้อมูลโอเพนซอร์สสองแบบ (การศึกษา B)"),
    # ---- 5.3 learning surface ----
    (", inspect every frame on the wire, and witness the security "
     "consequences of leaving those frames unauthenticated.",
     ", and inspect every frame on the wire."),
    ("ตรวจสอบทุกเฟรมบนสายสัญญาณ และเห็นผลกระทบด้านความปลอดภัยของการปล่อยให้เฟรมเหล่านั้น"
     "ไม่ผ่านการยืนยันตัวตน",
     "และตรวจสอบทุกเฟรมบนสายสัญญาณ"),
    # ---- 5.5 future work ----
    ("Several extensions would broaden the platform. On security, MQTT "
     "authentication and TLS, firewalling of Modbus TCP, and VLAN segregation "
     "of the OT and IT zones would convert the demonstrated weaknesses into "
     "enforced controls. On the gateway,",
     "Several extensions would broaden the platform. On the gateway,"),
    ("adding OPC-UA as a more secure OT protocol and validating",
     "adding OPC-UA and validating"),
    ("ส่วนขยายหลายส่วนสามารถขยายขอบเขตของแพลตฟอร์มได้ ด้านความปลอดภัย การยืนยันตัวตนและ TLS "
     "สำหรับ MQTT การจำกัด Modbus TCP ด้วยไฟร์วอลล์ และการแยกโซนโอทีและไอทีด้วย VLAN จะเปลี่ยน"
     "จุดอ่อนที่สาธิตไว้ให้กลายเป็นการควบคุมที่บังคับใช้จริง ด้านเกตเวย์",
     "ส่วนขยายหลายส่วนสามารถขยายขอบเขตของแพลตฟอร์มได้ ด้านเกตเวย์"),
    ("การเพิ่ม OPC-UA เป็นโพรโทคอลโอทีที่ปลอดภัยกว่า และการทดสอบ",
     "การเพิ่ม OPC-UA และการทดสอบ"),
    # ---- stray citations of the removed refs ----
    ("deterministic control [10]", "deterministic control"),
    ("ที่มีความแน่นอนเชิงเวลาแทน [10]", "ที่มีความแน่นอนเชิงเวลาแทน"),
    ("[1], [9], [10]", "[1]"),
    # ---- renumber 2.8 -> 2.7 ----
    ("2.8  Related Work and Research Gap", "2.7  Related Work and Research Gap"),
    ("2.8  งานวิจัยที่เกี่ยวข้องและช่องว่างงานวิจัย", "2.7  งานวิจัยที่เกี่ยวข้องและช่องว่างงานวิจัย"),
]

DELETE_REFS = ("[9]", "[10]")


def set_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def has_sectpr(p):
    ppr = p._p.find(qn('w:pPr'))
    return ppr is not None and ppr.find(qn('w:sectPr')) is not None


def process(path):
    d = docx.Document(path)
    n_del = n_rep = 0

    # 1. delete whole sections
    paras = list(d.paragraphs)
    i = 0
    while i < len(paras):
        if paras[i].text.strip() in DELETE_SECTIONS:
            j = i
            while j < len(paras):
                p = paras[j]
                if j > i and (p.style.name.startswith('Heading') or has_sectpr(p)):
                    break
                p._p.getparent().remove(p._p)
                n_del += 1
                j += 1
            paras = list(d.paragraphs)
            i = 0
            continue
        i += 1

    # 2. text replacements
    for p in d.paragraphs:
        t = p.text
        new = t
        for old, rep in REPLACE:
            if old in new:
                new = new.replace(old, rep)
        if new != t:
            set_text(p, new)
            n_rep += 1

    # 3. delete references [9], [10]
    for p in list(d.paragraphs):
        if p.style.name == 'Bibliography' and p.text.strip().startswith(DELETE_REFS):
            p._p.getparent().remove(p._p)
            n_del += 1

    d.save(path)
    print(f"{os.path.basename(path)}: removed {n_del} paragraphs, rewrote {n_rep}.")


if __name__ == '__main__':
    for name in ('thesis_en.docx', 'thesis_th.docx'):
        f = os.path.join(ROOT, name)
        if os.path.exists(f):
            process(f)
